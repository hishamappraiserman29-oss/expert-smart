"""
professional_valuation_core_excel_generator.py
Generates admin Excel workbooks for the three core Professional Valuation reports.
advisory_only=True | not_real_training=True
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from professional_valuation_core_report_examples import (
    SUBJECT, COMPARABLES, INCOME, COST, AVM, RECONCILIATION,
    DCF_5YEAR, DCF_10YEAR, SCENARIOS, SENSITIVITY_MATRIX, RISKS,
    STANDARDS, DATA_INPUTS, METHOD_SELECTION, ADVISORY_NOTE,
    RECONCILIATION_SCORECARD,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False

_GENERATED_AT = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

_NAVY  = "1A3A5C"
_GOLD  = "D4AF37"
_LIGHT = "E4EDF5"
_WARN  = "FFF3CD"


def _hdr(ws, row: int, col: int, value: str, bold: bool = True, bg: str = _NAVY, fg: str = "FFFFFF") -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Calibri", bold=bold, color=fg, size=10)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _cell(ws, row: int, col: int, value, bold: bool = False, italic: bool = False, bg: str = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Calibri", bold=bold, italic=italic, size=10)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _write_kv_block(ws, start_row: int, pairs: list[tuple], label_bg: str = _LIGHT) -> int:
    row = start_row
    for k, v in pairs:
        _cell(ws, row, 1, k, bold=True, bg=label_bg)
        _cell(ws, row, 2, str(v))
        row += 1
    return row


def _write_table(ws, start_row: int, headers: list[str], rows: list[list],
                 note: str = "") -> int:
    r = start_row
    for c, h in enumerate(headers, 1):
        _hdr(ws, r, c, h)
    r += 1
    for data_row in rows:
        bg = _LIGHT if (r - start_row - 1) % 2 == 0 else None
        for c, val in enumerate(data_row, 1):
            _cell(ws, r, c, str(val), bg=bg)
        r += 1
    if note:
        _cell(ws, r, 1, f"* {note}", italic=True)
        r += 1
    return r + 1


def _sheet_cover(wb: Workbook, report_type: str, label: str) -> None:
    ws = wb.create_sheet("Cover")
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 55

    _hdr(ws, 1, 1, label.upper(), bg=_NAVY)
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 30

    _hdr(ws, 2, 1, "ADVISORY ONLY — NOT FOR OFFICIAL USE", bg=_WARN, fg="8A5700")
    ws.merge_cells("A2:B2")

    pairs = [
        ("Report Type",    label),
        ("Report Key",     report_type),
        ("Request ID",     SUBJECT["request_id"]),
        ("Property",       SUBJECT["property_address"]),
        ("Property Type",  SUBJECT["property_type"]),
        ("GFA",            f'{SUBJECT["gross_floor_area_m2"]} m²'),
        ("Purpose",        SUBJECT["purpose"]),
        ("Basis of Value", SUBJECT["basis_of_value"]),
        ("Valuation Date", SUBJECT["valuation_date"]),
        ("Client",         SUBJECT["client"]),
        ("Currency",       SUBJECT["currency"]),
        ("Generated",      _GENERATED_AT),
        ("advisory_only",  "True"),
        ("fake_approval_created", "False"),
        ("certification_ready",   "False"),
        ("Advisory Note",  ADVISORY_NOTE),
    ]
    _write_kv_block(ws, 4, pairs)


def _sheet_data_quality(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Quality")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 55
    _hdr(ws, 1, 1, "Data Quality and Completeness", bg=_NAVY)
    ws.merge_cells("A1:B1")
    pairs = [
        ("Data Quality Level",     "Acceptable (Illustrative)"),
        ("Minimum Required Fields","Complete — address, type, area, purpose, basis"),
        ("Comparable Evidence",    "5 sales within 2 km"),
        ("Income Data",            "Indicative — market rent and vacancy"),
        ("Cost Data",              "Indicative — RCN and depreciation"),
        ("Legal / Title",          "Indicative — confirmation pending"),
        ("Report Readiness",       "Advisory Draft — Expert Review Required"),
        ("Missing Inputs",         "Title confirmation; lease data; inspection certificate"),
        ("expert_review_required", "True"),
    ]
    _write_kv_block(ws, 3, pairs)


def _sheet_property(wb: Workbook) -> None:
    ws = wb.create_sheet("Property")
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 55
    _hdr(ws, 1, 1, "Property Identification", bg=_NAVY)
    ws.merge_cells("A1:B1")
    pairs = [(k.replace("_", " ").title(), str(v)) for k, v in SUBJECT.items()]
    _write_kv_block(ws, 3, pairs)


def _sheet_comparables(wb: Workbook) -> None:
    ws = wb.create_sheet("Comparables")
    for i, w in enumerate([6, 35, 10, 12, 12, 12, 10, 10, 10, 10], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Comparable Sales (Indicative Only)")
    ws.merge_cells("A1:J1")
    headers = ["ID", "Address", "Area m²", "Total SAR", "SAR/m²", "Age Yrs", "Floor", "Parking", "Condition", "Sale Date"]
    _write_table(ws, 2, headers,
                 [[c["id"], c["address"], c["area_m2"], f'{c["price_sar"]:,}',
                   f'{c["price_per_sqm"]:,}', c["age_years"], c["floor"],
                   c["parking"], c["condition"], c["sale_date"]] for c in COMPARABLES],
                 note=ADVISORY_NOTE)


def _sheet_income(wb: Workbook) -> None:
    ws = wb.create_sheet("Income")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    _hdr(ws, 1, 1, "Income Approach — Operating Assumptions (Indicative)")
    ws.merge_cells("A1:C1")
    I = INCOME
    rows = [
        ["Gross Market Rent (Annual)", f'SAR {I["gross_market_rent_annual"]:,}', "100%"],
        [f'Less Vacancy ({I["vacancy_rate_pct"]}%)', f'(SAR {I["vacancy_deduction"]:,})', ""],
        ["Effective Gross Income", f'SAR {I["effective_gross_income"]:,}', ""],
        [f'Management ({I["management_fee_pct"]}%)', f'(SAR {I["management_fee"]:,})', ""],
        [f'Maintenance ({I["maintenance_pct"]}%)', f'(SAR {I["maintenance"]:,})', ""],
        [f'Insurance ({I["insurance_pct"]}%)', f'(SAR {I["insurance"]:,})', ""],
        [f'Service Charge ({I["service_charge_pct"]}%)', f'(SAR {I["service_charge"]:,})', ""],
        ["Other OpEx", f'(SAR {I["other_opex"]:,})', ""],
        ["Total OpEx", f'(SAR {I["total_opex"]:,})', ""],
        ["Net Operating Income (NOI)", f'SAR {I["net_operating_income"]:,}', ""],
        ["Capitalisation Rate", f'{I["cap_rate_pct"]}%', ""],
        ["Indicated Income Value", f'SAR {I["indicated_income_value"]:,}', ""],
    ]
    _write_table(ws, 3, ["Item", "Amount (SAR)", "Note"], rows, note=ADVISORY_NOTE)


def _sheet_cost(wb: Workbook) -> None:
    ws = wb.create_sheet("Cost")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    _hdr(ws, 1, 1, "Cost Approach — Inputs and Depreciation (Indicative)")
    ws.merge_cells("A1:C1")
    C = COST
    rows = [
        ["Land Area", f'{C["land_area_m2"]} m²', ""],
        ["Land Value Rate", f'SAR {C["land_value_per_sqm"]:,} / m²', "Indicative"],
        ["Land Value", f'SAR {C["land_value"]:,}', ""],
        ["RCN Rate", f'SAR {C["replacement_cost_per_sqm"]:,} / m²', "Indicative"],
        ["Gross Replacement Cost", f'SAR {C["gross_replacement_cost"]:,}', ""],
        ["Physical Depreciation %", f'{C["physical_depreciation_pct"]}%', "Age/life"],
        ["Depreciation SAR", f'SAR {C["total_depreciation_sar"]:,}', ""],
        ["Depreciated Improvement Value", f'SAR {C["depreciated_improvement_value"]:,}', ""],
        ["Indicated Cost Value", f'SAR {C["indicated_cost_value"]:,}', ""],
    ]
    _write_table(ws, 3, ["Item", "Value", "Note"], rows, note=ADVISORY_NOTE)


def _sheet_reconciliation(wb: Workbook) -> None:
    ws = wb.create_sheet("Reconciliation")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    _hdr(ws, 1, 1, "Reconciliation and Value Conclusion (Indicative)")
    ws.merge_cells("A1:E1")
    R = RECONCILIATION
    rows = [
        ["Market Approach", f'SAR {R["market_value"]:,}',  f'{R["market_approach_weight_pct"]}%',
         f'SAR {R["market_contribution"]:,}', "High"],
        ["Income Approach", f'SAR {R["income_value"]:,}',  f'{R["income_approach_weight_pct"]}%',
         f'SAR {R["income_contribution"]:,}', "Medium"],
        ["Cost Approach",   f'SAR {R["cost_value"]:,}',   f'{R["cost_approach_weight_pct"]}%',
         f'SAR {R["cost_contribution"]:,}', "Medium"],
        ["Weighted Total",  "—", "100%", f'SAR {R["weighted_total"]:,}', ""],
        ["Adopted Value",   f'SAR {R["adopted_value"]:,}', "—", "—", "Rounded"],
    ]
    _write_table(ws, 3, ["Approach", "Indicated Value", "Weight", "Contribution", "Reliability"], rows,
                 note=ADVISORY_NOTE)
    _cell(ws, ws.max_row + 1, 1, "Rationale:", bold=True)
    _cell(ws, ws.max_row, 2, R["rationale"])


def _sheet_admin_notes(wb: Workbook, report_type: str) -> None:
    ws = wb.create_sheet("Admin Notes")
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 60
    _hdr(ws, 1, 1, "Admin / Expert Review Notes", bg=_GOLD, fg="1A1A1A")
    ws.merge_cells("A1:B1")
    pairs = [
        ("advisory_only",              "True"),
        ("not_real_training",          "True"),
        ("fake_approval_created",      "False"),
        ("fake_signature_created",     "False"),
        ("certification_ready",        "False"),
        ("expert_review_required",     "True"),
        ("report_type",                report_type),
        ("generated_at",               _GENERATED_AT),
        ("data_quality_level",         "Acceptable (Illustrative)"),
        ("minimum_requirements_met",   "Partially — see Data Quality sheet"),
        ("recommended_actions",        "Submit to licensed valuation expert for review and certification"),
        ("internal_paths_exposed",     "False"),
        ("advisory_note",              ADVISORY_NOTE),
    ]
    _write_kv_block(ws, 3, pairs, label_bg="E4EDF5")


def _sheet_dcf(wb: Workbook, use_10yr: bool = False) -> None:
    ws = wb.create_sheet("DCF")
    D = DCF_10YEAR if use_10yr else DCF_5YEAR
    yrs = len(D["years"])
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    label = f'DCF {"10-Year" if use_10yr else "5-Year"} Projection (Indicative)'
    _hdr(ws, 1, 1, label)
    ws.merge_cells(f"A1:D1")
    _hdr(ws, 2, 1, f'Discount Rate: {D["discount_rate_pct"]}% | Growth: {D["growth_rate_pct"]}% | Terminal Cap: {D["terminal_cap_rate_pct"]}%',
         bg=_LIGHT, fg="1A1A1A")
    ws.merge_cells(f"A2:D2")
    _write_table(ws, 3, ["Year", "NOI (SAR)", "PV Factor", "PV of NOI"],
                 [[D["years"][i], f'SAR {D["noi"][i]:,}', f'{D["pv_factor"][i]:.4f}',
                   f'SAR {D["pv_noi"][i]:,}'] for i in range(yrs)])
    r = ws.max_row + 1
    for label_t, val in [
        ("Cumulative PV of NOI", f'SAR {D["cumulative_pv_noi"]:,}'),
        ("Terminal Value",       f'SAR {D["terminal_value"]:,}'),
        ("PV of Terminal Value", f'SAR {D["pv_terminal"]:,}'),
        ("Total DCF Value",      f'SAR {D["total_dcf_value"]:,}'),
    ]:
        _cell(ws, r, 1, label_t, bold=True, bg=_LIGHT)
        _cell(ws, r, 2, val, bold=True)
        r += 1
    _cell(ws, r + 1, 1, f"* {ADVISORY_NOTE}", italic=True)


def _sheet_scenarios(wb: Workbook) -> None:
    ws = wb.create_sheet("Scenarios")
    for i, w in enumerate([16, 10, 14, 10, 12, 14, 14, 14, 12], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Scenario Analysis (Indicative)")
    ws.merge_cells("A1:I1")
    headers = ["Scenario", "Growth %", "Disc Rate %", "Cap Rate %", "Occupancy %",
               "DCF Value", "Market Value", "Adopted Value", "Probability %"]
    rows = [[s["name"], f'{s["growth_pct"]}%', f'{s["discount_rate_pct"]}%',
             f'{s["cap_rate_pct"]}%', f'{s["occupancy_pct"]}%',
             f'SAR {s["dcf_value"]:,}', f'SAR {s["market_value"]:,}',
             f'SAR {s["adopted_value"]:,}', f'{s["probability_pct"]}%'] for s in SCENARIOS]
    from professional_valuation_core_report_examples import PROBABILITY_WEIGHTED_VALUE
    rows.append(["Prob-Weighted", "—", "—", "—", "—", "—", "—",
                 f'SAR {PROBABILITY_WEIGHTED_VALUE:,}', "100%"])
    _write_table(ws, 3, headers, rows, note=ADVISORY_NOTE)


def _sheet_risks(wb: Workbook) -> None:
    ws = wb.create_sheet("Risk Register")
    for i, w in enumerate([5, 16, 38, 12, 10, 8, 38], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Risk Register (Indicative)")
    ws.merge_cells("A1:G1")
    rows = [[r["id"], r["category"], r["description"], r["probability"],
             r["impact"], r["score"], r["mitigation"]] for r in RISKS]
    _write_table(ws, 3, ["ID", "Category", "Description", "Probability", "Impact", "Score", "Mitigation"],
                 rows, note=ADVISORY_NOTE)


def _sheet_standards(wb: Workbook) -> None:
    ws = wb.create_sheet("Standards Matrix")
    for i, w in enumerate([12, 30, 18, 38], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Applied Standards Matrix (Indicative)")
    ws.merge_cells("A1:D1")
    rows = [[s["code"], s["name"], s["status"], s["notes"]] for s in STANDARDS]
    _write_table(ws, 3, ["Code", "Standard", "Status", "Notes"], rows,
                 note="Indicative compliance mapping — expert confirmation required")


def _sheet_method_selection(wb: Workbook) -> None:
    ws = wb.create_sheet("Method Selection")
    for i, w in enumerate([32, 14, 12, 10, 40], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Valuation Method Selection Matrix")
    ws.merge_cells("A1:E1")
    rows = [[m["method"], m["data_available"], m["applicable"],
             f'{m["weight_pct"]}%', m["justification"]] for m in METHOD_SELECTION]
    _write_table(ws, 3, ["Method", "Data Available", "Applicable", "Weight", "Justification"], rows)


def _sheet_weighted_reconciliation(wb: Workbook) -> None:
    ws = wb.create_sheet("Wtd Reconciliation")
    for i, w in enumerate([24, 18, 10, 12, 14, 18], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    _hdr(ws, 1, 1, "Weighted Reconciliation Scorecard (Indicative)")
    ws.merge_cells("A1:F1")
    rows = []
    for r in RECONCILIATION_SCORECARD:
        iv = f'SAR {r["indicated_value"]:,}' if r["indicated_value"] else "—"
        w  = f'{r["weight_pct"]}%' if r["weight_pct"] is not None else "—"
        c  = f'SAR {r["contribution"]:,}' if r["contribution"] else "—"
        rows.append([r["method"], iv, w, r["reliability"], r["data_quality"], c])
    _write_table(ws, 3, ["Method", "Indicated Value", "Weight", "Reliability", "Data Quality", "Contribution"],
                 rows, note=ADVISORY_NOTE)


# ── Public generators ─────────────────────────────────────────────────────────

def generate_traditional_workbook(path: Path) -> bool:
    if not _OPENPYXL_OK:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    _sheet_cover(wb, "traditional_report", "Traditional Valuation Report")
    _sheet_data_quality(wb)
    _sheet_property(wb)
    _sheet_comparables(wb)
    _sheet_income(wb)
    _sheet_cost(wb)
    _sheet_reconciliation(wb)
    _sheet_admin_notes(wb, "traditional_report")
    wb.save(str(path))
    return path.exists() and path.stat().st_size > 1000


def generate_detailed_workbook(path: Path) -> bool:
    if not _OPENPYXL_OK:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    _sheet_cover(wb, "detailed_report", "Detailed Valuation Report")
    _sheet_data_quality(wb)
    _sheet_property(wb)
    _sheet_comparables(wb)
    _sheet_income(wb)
    _sheet_cost(wb)
    _sheet_dcf(wb, use_10yr=False)
    _sheet_reconciliation(wb)
    _sheet_risks(wb)
    _sheet_admin_notes(wb, "detailed_report")
    wb.save(str(path))
    return path.exists() and path.stat().st_size > 1000


def generate_professional_workbook(path: Path) -> bool:
    if not _OPENPYXL_OK:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    _sheet_cover(wb, "professional_report", "Professional Valuation Report")
    _sheet_data_quality(wb)
    _sheet_property(wb)
    _sheet_comparables(wb)
    _sheet_income(wb)
    _sheet_cost(wb)
    _sheet_dcf(wb, use_10yr=True)
    _sheet_scenarios(wb)
    _sheet_method_selection(wb)
    _sheet_reconciliation(wb)
    _sheet_weighted_reconciliation(wb)
    _sheet_risks(wb)
    _sheet_standards(wb)
    _sheet_admin_notes(wb, "professional_report")
    wb.save(str(path))
    return path.exists() and path.stat().st_size > 1000


def generate_all_excel(excel_dir: Path) -> dict:
    excel_dir.mkdir(parents=True, exist_ok=True)
    if not _OPENPYXL_OK:
        return {
            "traditional_report": {"status": "BLOCKED", "reason": "openpyxl not installed", "exists": False},
            "detailed_report":    {"status": "BLOCKED", "reason": "openpyxl not installed", "exists": False},
            "professional_report":{"status": "BLOCKED", "reason": "openpyxl not installed", "exists": False},
        }

    results = {}
    specs = [
        ("traditional_report",  "traditional_report_admin_workbook.xlsx",  generate_traditional_workbook),
        ("detailed_report",     "detailed_report_admin_workbook.xlsx",     generate_detailed_workbook),
        ("professional_report", "professional_report_admin_workbook.xlsx", generate_professional_workbook),
    ]
    for key, fname, gen_fn in specs:
        path = excel_dir / fname
        ok = gen_fn(path)
        results[key] = {
            "status":       "OK" if ok else "FAILED",
            "file_name":    fname,
            "file_path":    str(path),
            "exists":       path.exists(),
            "size_bytes":   path.stat().st_size if path.exists() else 0,
            "blocker":      None if ok else "Generation failed",
        }
    return results


if __name__ == "__main__":
    here = Path(__file__).parent
    qa_root = here / "instance" / "manual_review_outputs" / \
              "professional_valuation_final_core_workflow_and_report_qa"
    res = generate_all_excel(qa_root / "excel_outputs")
    for k, v in res.items():
        print(f"[{v['status']}] {k}: {v.get('size_bytes', 0):,} bytes")
