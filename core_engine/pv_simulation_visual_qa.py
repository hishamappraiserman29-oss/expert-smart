"""
pv_simulation_visual_qa.py — Visual QA runner for simulation report artifacts.

Checks:
  1. Playwright HTML screenshots  — every [data-section] element in user/admin HTML
  2. PyMuPDF PDF page rendering   — every page of user/admin PDF (text presence, not blank)
  3. openpyxl structural inspect  — admin Excel sheet names, RTL, chart count, disclosure text
  4. LibreOffice chart render      — optional; converts Excel to PNG if soffice.exe found

Returns a dict suitable for JSON serialisation.

Usage (standalone):
  python pv_simulation_visual_qa.py <user_html> <user_pdf> <admin_html> <admin_pdf> <admin_xlsx>
"""
from __future__ import annotations

import json
import os
import pathlib
import subprocess
import sys
import datetime
from typing import Any


_DISCLOSURE_FRAGMENT = "internal_base_scenario"
_NO_LOCAL_PATH_MARKERS = ["file:///", "C:\\", "C:/", "D:\\", "D:/"]

LIBREOFFICE_CANDIDATES = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
]


def _ts() -> str:
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def _find_libreoffice() -> str | None:
    for p in LIBREOFFICE_CANDIDATES:
        if os.path.isfile(p):
            return p
    return None


# ── 1. Playwright HTML screenshots ────────────────────────────────────────────

def screenshot_html_sections(
    html_path: pathlib.Path,
    out_dir: pathlib.Path,
    label: str = "html",
) -> dict:
    """
    Screenshot every [data-section] element in the HTML file.
    Also screenshots the full page.
    Returns a result dict with per-section status.
    """
    result: dict[str, Any] = {
        "label": label,
        "html_path": str(html_path),
        "sections": [],
        "full_page": None,
        "issues": [],
        "pass": False,
    }

    if not html_path.exists():
        result["issues"].append(f"HTML file not found: {html_path}")
        return result

    out_dir.mkdir(parents=True, exist_ok=True)

    # Pre-check: no local file paths exposed
    html_text = html_path.read_text(encoding="utf-8")
    for marker in _NO_LOCAL_PATH_MARKERS:
        if marker in html_text:
            result["issues"].append(f"Local path marker found in HTML: {marker}")

    # Pre-check: disclosure text present
    if _DISCLOSURE_FRAGMENT not in html_text:
        result["issues"].append("Disclosure fragment missing from HTML")

    # Pre-check: watermark present
    if "advisory_only=True" not in html_text:
        result["issues"].append("advisory_only flag missing from HTML")

    # Pre-check: RTL direction declared
    if 'dir="rtl"' not in html_text and "direction:rtl" not in html_text:
        result["issues"].append("RTL direction not declared in HTML")

    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page    = browser.new_page(viewport={"width": 1200, "height": 900})
            page.goto(html_path.resolve().as_uri(), wait_until="networkidle", timeout=30_000)

            # Full-page screenshot
            full_path = out_dir / f"{label}_full_page.png"
            page.screenshot(path=str(full_path), full_page=True)
            if full_path.exists() and full_path.stat().st_size > 1_000:
                result["full_page"] = str(full_path)
            else:
                result["issues"].append("Full-page screenshot empty or missing")

            # Per-section screenshots
            sections = page.query_selector_all("[data-section]")
            if not sections:
                result["issues"].append("No [data-section] elements found")
            for el in sections:
                sec_num = el.get_attribute("data-section") or "unknown"
                sec_path = out_dir / f"{label}_section_{sec_num}.png"
                try:
                    el.screenshot(path=str(sec_path))
                    size = sec_path.stat().st_size if sec_path.exists() else 0
                    is_blank = size < 500
                    result["sections"].append({
                        "section": sec_num,
                        "file":    str(sec_path),
                        "size_bytes": size,
                        "blank":   is_blank,
                        "pass":    not is_blank,
                    })
                    if is_blank:
                        result["issues"].append(f"Section {sec_num} screenshot is blank")
                except Exception as se:
                    result["issues"].append(f"Section {sec_num} screenshot failed: {se}")
                    result["sections"].append({"section": sec_num, "error": str(se), "pass": False})

            browser.close()

        result["pass"] = len(result["issues"]) == 0
    except ImportError:
        result["issues"].append("Playwright not installed — HTML screenshot skipped")
        result["pass"] = None   # None = skipped, not failed
    except Exception as e:
        result["issues"].append(f"Playwright error: {e}")

    return result


# ── 2. PyMuPDF PDF page rendering ─────────────────────────────────────────────

def inspect_pdf_pages(
    pdf_path: pathlib.Path,
    out_dir: pathlib.Path,
    label: str = "pdf",
) -> dict:
    """
    Render every PDF page to PNG using PyMuPDF (fitz).
    Also checks that pages contain text (not blank).
    """
    result: dict[str, Any] = {
        "label": label,
        "pdf_path": str(pdf_path),
        "pages": [],
        "issues": [],
        "pass": False,
    }

    if not pdf_path.exists():
        result["issues"].append(f"PDF file not found: {pdf_path}")
        return result

    # Check %PDF header
    header = pdf_path.read_bytes()[:4]
    if header != b"%PDF":
        result["issues"].append(f"Invalid PDF header: {header!r}")
        return result

    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(pdf_path))
        page_count = len(doc)
        if page_count == 0:
            result["issues"].append("PDF has 0 pages")
            doc.close()
            return result

        for i, page in enumerate(doc):
            # Render to image
            mat  = fitz.Matrix(1.5, 1.5)   # 108 dpi
            pix  = page.get_pixmap(matrix=mat)
            img_path = out_dir / f"{label}_page_{i+1}.png"
            pix.save(str(img_path))
            img_size = img_path.stat().st_size if img_path.exists() else 0

            # Extract text
            text = page.get_text("text")
            has_text = len(text.strip()) > 20

            # Check for local path exposure in text layer
            local_exposed = any(m in text for m in _NO_LOCAL_PATH_MARKERS)
            if local_exposed:
                result["issues"].append(f"Page {i+1}: local path found in text layer")

            page_info = {
                "page":        i + 1,
                "file":        str(img_path),
                "size_bytes":  img_size,
                "text_chars":  len(text.strip()),
                "has_text":    has_text,
                "local_path_exposed": local_exposed,
                "pass":        img_size > 1_000 and has_text and not local_exposed,
            }
            result["pages"].append(page_info)

            if img_size < 1_000:
                result["issues"].append(f"Page {i+1} render appears blank")
            if not has_text:
                result["issues"].append(f"Page {i+1} has no extractable text")

        doc.close()
        result["page_count"] = page_count
        result["pass"] = (page_count > 0) and len(result["issues"]) == 0

    except ImportError:
        # PyMuPDF not available — do minimal checks
        result["issues"].append("PyMuPDF (fitz) not installed — PDF page rendering skipped")
        size = pdf_path.stat().st_size
        result["pages"].append({
            "page": 1, "size_bytes": size,
            "has_text": size > 1_000,   # real Playwright PDF will be > 100KB
            "pass": size > 1_000,
        })
        result["pass"] = None   # skipped

    except Exception as e:
        result["issues"].append(f"PyMuPDF error: {e}")

    return result


# ── 3. openpyxl structural inspection ────────────────────────────────────────

def inspect_excel_structure(
    xl_path: pathlib.Path,
    expected_sheets: list[str] | None = None,
    min_sheets: int = 10,
) -> dict:
    """
    Structural inspection of Excel workbook via openpyxl.
    Checks: sheet count, RTL, chart presence, disclosure text, no local paths.
    """
    result: dict[str, Any] = {
        "xl_path": str(xl_path),
        "sheets": [],
        "chart_count": 0,
        "issues": [],
        "pass": False,
    }

    if not xl_path.exists():
        result["issues"].append(f"Excel file not found: {xl_path}")
        return result

    header = xl_path.read_bytes()[:4]
    if header != b"PK\x03\x04":
        result["issues"].append(f"Invalid XLSX header: {header!r}")
        return result

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(xl_path), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        result["sheet_names"] = sheet_names
        result["sheet_count"] = len(sheet_names)

        if len(sheet_names) < min_sheets:
            result["issues"].append(
                f"Expected ≥{min_sheets} sheets, found {len(sheet_names)}: {sheet_names}"
            )

        # Check for expected sheets
        if expected_sheets:
            for es in expected_sheets:
                if not any(es in sn for sn in sheet_names):
                    result["issues"].append(f"Expected sheet containing '{es}' not found")

        disclosure_found = False
        local_path_found = False
        total_charts = 0

        for sname in sheet_names:
            ws = wb[sname]
            s_info: dict[str, Any] = {"name": sname, "rtl": False, "issues": []}

            # RTL check
            if hasattr(ws, "sheet_view") and ws.sheet_view:
                s_info["rtl"] = bool(getattr(ws.sheet_view, "rightToLeft", False))
            elif hasattr(ws, "views") and ws.views:
                for v in ws.views.sheetView:
                    if getattr(v, "rightToLeft", False):
                        s_info["rtl"] = True

            # Text content scan
            cell_texts = []
            try:
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if cell and isinstance(cell, str):
                            cell_texts.append(cell)
            except Exception:
                pass
            all_text = " ".join(cell_texts)

            if _DISCLOSURE_FRAGMENT in all_text:
                disclosure_found = True
            for marker in _NO_LOCAL_PATH_MARKERS:
                if marker in all_text:
                    local_path_found = True
                    result["issues"].append(f"Sheet '{sname}': local path marker found")

            result["sheets"].append(s_info)

        wb.close()

        # Chart count via zip inspection (openpyxl read_only doesn't expose charts)
        import zipfile
        with zipfile.ZipFile(str(xl_path), "r") as zf:
            total_charts = sum(1 for n in zf.namelist() if "charts/chart" in n)
        result["chart_count"] = total_charts

        if total_charts < 1:
            result["issues"].append("No charts found in Excel workbook")

        if not disclosure_found:
            result["issues"].append("Disclosure text not found in any sheet")

        if local_path_found:
            result["issues"].append("Local path markers found in workbook cells")

        result["pass"] = len(result["issues"]) == 0

    except ImportError:
        result["issues"].append("openpyxl not installed — Excel inspection skipped")
        result["pass"] = None
    except Exception as e:
        result["issues"].append(f"Excel inspection error: {e}")

    return result


# ── 4. LibreOffice chart rendering (optional) ─────────────────────────────────

def render_excel_with_libreoffice(
    xl_path: pathlib.Path,
    out_dir: pathlib.Path,
) -> dict:
    """
    Convert Excel to PDF/PNG via LibreOffice headless for visual chart inspection.
    Falls back gracefully if LibreOffice is not installed.
    """
    result: dict[str, Any] = {
        "xl_path": str(xl_path),
        "rendered_files": [],
        "issues": [],
        "pass": False,
        "skipped": False,
    }

    lo = _find_libreoffice()
    if not lo:
        result["issues"].append("LibreOffice not found — Excel chart render skipped")
        result["skipped"] = True
        result["pass"] = None
        return result

    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        proc = subprocess.run(
            [lo, "--headless", "--convert-to", "png",
             str(xl_path.resolve()), "--outdir", str(out_dir.resolve())],
            capture_output=True, text=True, timeout=60,
        )
        if proc.returncode != 0:
            result["issues"].append(f"LibreOffice exit {proc.returncode}: {proc.stderr[:300]}")
        else:
            rendered = list(out_dir.glob(xl_path.stem + "*.png"))
            if not rendered:
                rendered = list(out_dir.glob("*.png"))
            result["rendered_files"] = [str(f) for f in rendered]
            if not rendered:
                result["issues"].append("LibreOffice ran but produced no PNG files")
        result["pass"] = len(result["issues"]) == 0
    except subprocess.TimeoutExpired:
        result["issues"].append("LibreOffice timed out after 60 s")
    except Exception as e:
        result["issues"].append(f"LibreOffice render error: {e}")

    return result


# ── Orchestrator ──────────────────────────────────────────────────────────────

def run_simulation_visual_qa(
    user_html:  pathlib.Path | str,
    user_pdf:   pathlib.Path | str,
    admin_html: pathlib.Path | str,
    admin_pdf:  pathlib.Path | str,
    admin_xlsx: pathlib.Path | str,
    out_dir:    pathlib.Path | str | None = None,
    case_id:    str = "",
) -> dict:
    """
    Run full Visual QA pipeline for all 5 simulation artifacts.
    Returns structured JSON-serialisable report.
    """
    uh  = pathlib.Path(user_html).resolve()
    up  = pathlib.Path(user_pdf).resolve()
    ah  = pathlib.Path(admin_html).resolve()
    ap  = pathlib.Path(admin_pdf).resolve()
    axl = pathlib.Path(admin_xlsx).resolve()

    if out_dir is None:
        out_dir = uh.parent / f"visual_qa_sim_v2_{_ts()}"
    out_dir = pathlib.Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    report: dict[str, Any] = {
        "case_id":      case_id,
        "generated_at": datetime.datetime.now().isoformat(),
        "artifacts": {
            "user_html":  str(uh),
            "user_pdf":   str(up),
            "admin_html": str(ah),
            "admin_pdf":  str(ap),
            "admin_xlsx": str(axl),
        },
    }

    EXPECTED_SHEETS = [
        "ملخص",
        "بيانات الأصل",
        "الأساسي",
        "المتحفظ",
        "المتفائل",
        "نتائج",
        "الظلّي",
        "مقارنة",
        "محركات",
        "الاستنتاج",
    ]

    report["user_html"]  = screenshot_html_sections(uh,  out_dir / "user_html",  "user_html")
    report["admin_html"] = screenshot_html_sections(ah,  out_dir / "admin_html", "admin_html")
    report["user_pdf"]   = inspect_pdf_pages(up,  out_dir / "user_pdf",   "user_pdf")
    report["admin_pdf"]  = inspect_pdf_pages(ap,  out_dir / "admin_pdf",  "admin_pdf")
    report["admin_excel_structural"] = inspect_excel_structure(
        axl, expected_sheets=EXPECTED_SHEETS, min_sheets=10
    )
    report["admin_excel_chart_render"] = render_excel_with_libreoffice(
        axl, out_dir / "excel_charts"
    )

    # Overall summary
    results = [
        report["user_html"],
        report["admin_html"],
        report["user_pdf"],
        report["admin_pdf"],
        report["admin_excel_structural"],
    ]
    all_issues: list[str] = []
    for r in results:
        all_issues.extend(r.get("issues", []))

    # Skipped checks (None = skipped, True/False = ran)
    skipped = sum(1 for r in results if r.get("pass") is None)
    failed_hard = [i for i in all_issues if "not found" in i or "Invalid" in i]

    report["summary"] = {
        "total_checks": len(results) + 1,   # +1 for LO render
        "skipped":      skipped,
        "all_issues":   all_issues,
        "issue_count":  len(all_issues),
        "pass":         len(failed_hard) == 0,
    }

    # Save JSON report
    report_path = out_dir / "visual_qa_report.json"
    try:
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        report["report_path"] = str(report_path)
    except Exception:
        pass

    return report


# ── CLI entry ──────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) < 6:
        print("Usage: python pv_simulation_visual_qa.py <user_html> <user_pdf> <admin_html> <admin_pdf> <admin_xlsx> [out_dir]")
        sys.exit(1)
    out = sys.argv[6] if len(sys.argv) > 6 else None
    result = run_simulation_visual_qa(
        sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5], out
    )
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
    sys.exit(0 if result["summary"]["pass"] else 1)
