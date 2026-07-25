"""
pv_sim_enhanced_excel.py
Adds two admin-only sheets + two BarCharts to the simulation Excel workbook:

  Sheet "المصدرية"      — Provenance table (all inputs + source/confidence/status)
  Sheet "مقارنة_مصدرية" — Methods comparison (computed vs simulation) + reconciliation

Charts (added to each sheet):
  - Bar chart: computed method values vs. simulation (Base) values
  - Bar chart: input source distribution (mass appraisal / web Draft / unavailable)

Admin only — the caller is responsible for never serving this file to external users.
"""
from __future__ import annotations

from typing import Any, Dict, List, Optional

from .pv_sim_sourced_inputs import InputProvenance, SourcingResult
from .pv_sim_market_methods import MethodsResult


# ── Style helpers ─────────────────────────────────────────────────────────────

def _rtl(ws: Any) -> None:
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass


def _hdr_fill():
    try:
        from openpyxl.styles import PatternFill
        return PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    except Exception:
        return None


def _hdr_font():
    try:
        from openpyxl.styles import Font
        return Font(bold=True, color="E2E8F0", size=10)
    except Exception:
        return None


def _bold_font():
    try:
        from openpyxl.styles import Font
        return Font(bold=True)
    except Exception:
        return None


def _set_header_row(ws: Any, headers: List[str]) -> None:
    fill = _hdr_fill()
    font = _hdr_font()
    for col_i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_i, value=h)
        if fill:
            try:
                c.fill = fill
            except Exception:
                pass
        if font:
            try:
                c.font = font
            except Exception:
                pass


def _col_widths(ws: Any, widths: List[int]) -> None:
    for col_i, w in enumerate(widths, 1):
        try:
            ws.column_dimensions[ws.cell(1, col_i).column_letter].width = w
        except Exception:
            pass


def _scrub_uri(uri: str) -> str:
    """Remove local paths — never expose internal file system paths."""
    if not uri:
        return ""
    if (uri.startswith("file://") or ":\\" in uri
            or (uri.startswith("/") and not uri.startswith("/api"))):
        return "[مسار داخلي — محجوب]"
    return uri


# ── Sheet 1: Provenance ───────────────────────────────────────────────────────

def add_provenance_sheet(wb: Any, provenance_table: List[InputProvenance]) -> None:
    name = "المصدرية"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(title=name)
    _rtl(ws)

    headers = [
        "معرّف المدخل", "المدخل (عربي)", "القيمة المستخدمة", "الوحدة",
        "نوع المصدر", "اسم المصدر", "URI",
        "تاريخ الاسترداد", "الثقة %", "الفئة", "الحالة", "ملاحظة التسوية",
        "قيمة التقييم المجمع", "قيمة البحث",
    ]
    _set_header_row(ws, headers)

    for row_i, p in enumerate(provenance_table, 2):
        row_data = [
            p.input_id, p.label_ar,
            p.value_used, p.unit,
            p.source_type, p.source_name, _scrub_uri(p.source_uri),
            p.retrieved_at[:10], p.confidence_score,
            p.source_tier, p.status, p.reconciliation_note,
            p.mass_appraisal_value, p.research_value,
        ]
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=row_i, column=col_i, value=val)

    _col_widths(ws, [18, 30, 18, 10, 18, 30, 40, 18, 10, 18, 30, 50, 20, 20])

    # Source distribution chart below the data
    _add_source_dist_chart(ws, provenance_table, data_row=len(provenance_table) + 4)


# ── Sheet 2: Methods Comparison ───────────────────────────────────────────────

def add_comparison_sheet(wb: Any, methods_result: MethodsResult) -> None:
    name = "مقارنة_مصدرية"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(title=name)
    _rtl(ws)

    headers = [
        "الأسلوب",
        "القيمة المحسوبة (الموضوع)",
        "قيمة التقرير (Base)",
        "الفرق",
        "الفرق %",
        "الحالة RAG",
        "مصدر المدخلات",
    ]
    _set_header_row(ws, headers)

    for row_i, row in enumerate(methods_result.comparison_table, 2):
        row_data = [
            row.get("method_name_ar") or row.get("method_id"),
            row.get("computed_value"),
            row.get("simulation_value"),
            row.get("diff_amount"),
            row.get("diff_pct"),
            row.get("rag_status"),
            row.get("input_source_summary"),
        ]
        for col_i, val in enumerate(row_data, 1):
            ws.cell(row=row_i, column=col_i, value=val)

    # Reconciliation row
    rec = methods_result.reconciliation
    rec_row = len(methods_result.comparison_table) + 3
    bf = _bold_font()
    for col_i, val in enumerate([
        "التوفيق الاستشاري المجمَّع",
        rec.get("reconciled_value"),
        None,
        rec.get("diff_from_simulation"),
        rec.get("diff_pct"),
        rec.get("rag_status"),
        rec.get("note"),
    ], 1):
        c = ws.cell(row=rec_row, column=col_i, value=val)
        if bf and col_i <= 2:
            try:
                c.font = bf
            except Exception:
                pass

    _col_widths(ws, [32, 24, 24, 18, 12, 10, 60])

    # Methods comparison chart
    chart_row = len(methods_result.comparison_table) + 6
    _add_methods_chart(ws, methods_result, chart_row)


# ── Chart helpers ─────────────────────────────────────────────────────────────

def _add_methods_chart(ws: Any, methods_result: MethodsResult, chart_row: int) -> None:
    """Bar chart: computed vs simulation per method (cols B and C)."""
    try:
        from openpyxl.chart import BarChart, Reference

        n = len(methods_result.comparison_table)
        if n == 0:
            return

        chart = BarChart()
        chart.type    = "col"
        chart.title   = "مقارنة قيم الأساليب: محسوب على الموضوع vs قيمة التقرير"
        chart.y_axis.title = methods_result.currency
        chart.x_axis.title = "الأسلوب"
        chart.style   = 10
        chart.width   = 22
        chart.height  = 14

        computed_ref  = Reference(ws, min_col=2, min_row=1, max_row=n + 1)
        sim_ref       = Reference(ws, min_col=3, min_row=1, max_row=n + 1)
        labels_ref    = Reference(ws, min_col=1, min_row=2, max_row=n + 1)

        chart.add_data(computed_ref, titles_from_data=True)
        chart.add_data(sim_ref,      titles_from_data=True)
        chart.set_categories(labels_ref)
        chart.series[0].title.v = "محسوب على الموضوع"
        chart.series[1].title.v = "قيمة التقرير"

        ws.add_chart(chart, f"A{chart_row}")
    except Exception:
        pass


def _add_source_dist_chart(
    ws: Any,
    provenance_table: List[InputProvenance],
    data_row: int,
) -> None:
    """Bar chart: input source distribution counts."""
    try:
        from openpyxl.chart import BarChart, Reference

        counts = {"نموذج مجمع (Certified)": 0, "بحث مسحي (Draft)": 0, "غير متاح": 0}
        for p in provenance_table:
            if p.source_type == "mass_appraisal":
                counts["نموذج مجمع (Certified)"] += 1
            elif p.source_type in ("web_research", "enrichment_layer"):
                counts["بحث مسحي (Draft)"] += 1
            else:
                counts["غير متاح"] += 1

        # Write distribution data
        ws.cell(row=data_row,     column=1, value="توزيع مصادر المدخلات")
        ws.cell(row=data_row,     column=2, value="عدد المدخلات")
        for i, (label, count) in enumerate(counts.items(), 1):
            ws.cell(row=data_row + i, column=1, value=label)
            ws.cell(row=data_row + i, column=2, value=count)

        chart = BarChart()
        chart.type   = "col"
        chart.title  = "توزيع مصادر المدخلات"
        chart.y_axis.title = "عدد المدخلات"
        chart.style  = 10
        chart.width  = 14
        chart.height = 10

        data_ref   = Reference(ws, min_col=2, min_row=data_row,     max_row=data_row + 3)
        labels_ref = Reference(ws, min_col=1, min_row=data_row + 1, max_row=data_row + 3)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(labels_ref)

        ws.add_chart(chart, f"D{data_row}")
    except Exception:
        pass


# ── Public API ────────────────────────────────────────────────────────────────

def add_enhanced_sheets(
    wb: Any,
    provenance_table: List[InputProvenance],
    methods_result: MethodsResult,
) -> Any:
    """
    Add Provenance + Comparison sheets (+ charts) to existing workbook.
    Returns the modified workbook.
    ADMIN ONLY — caller must never serve this file to external users.
    """
    add_provenance_sheet(wb, provenance_table)
    add_comparison_sheet(wb, methods_result)
    return wb
