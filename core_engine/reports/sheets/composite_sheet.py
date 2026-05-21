"""
composite_sheet.py — Composite property valuation Excel sheet (Wave 6).

Appends the "التقييم المركب" tab to an existing openpyxl Workbook.
Styling follows the Midnight Gold theme (report_theme.py).
"""
from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from reports.report_theme import (
    NumFormat,
    Palette,
    apply_sheet_defaults,
    draw_banner,
    draw_section,
    get_alignment,
    get_fill,
    get_font,
    style_body,
    style_table_header,
    style_value,
)

if TYPE_CHECKING:
    from adapters.purpose_adapter import AdjustedValuation
    from validation.composite_rules import CompositeValidationReport

_SHEET_NAME = "التقييم المركب"
_END_COL = 9

_HEADERS = [
    "م",
    "المعرّف",
    "الاسم",
    "نوع الأصل",
    "الغرض",
    "المساحة م²",
    "سعر القاعدة EGP/م²",
    "القيمة الأساسية EGP",
    "القيمة المعدّلة EGP",
]

_COL_WIDTHS = [5, 12, 22, 28, 42, 12, 18, 20, 20]


def build_composite_sheet(
    wb: Workbook,
    valuations: list[AdjustedValuation],
    *,
    raw_components: list[dict] | None = None,
    validation: CompositeValidationReport | None = None,
) -> None:
    """Append the composite valuation sheet to wb.

    Args:
        wb: Target workbook (not yet saved).
        valuations: Per-component adjusted valuation results (Wave 2 output).
        raw_components: Optional raw component dicts from the API request body;
            used to surface area_sqm and base_rate_per_sqm, which are not
            carried by AdjustedValuation.
        validation: Optional Wave 3 validation report; issues rendered in a
            second section when present and non-empty.
    """
    ws = wb.create_sheet(_SHEET_NAME)
    apply_sheet_defaults(ws, freeze=False)
    ws.sheet_view.rightToLeft = True

    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Row 1: Banner ─────────────────────────────────────────────────
    draw_banner(ws, row=1, end_col=_END_COL, text="التقييم المركب — نتائج التقييم الإجمالي")

    # ── Row 2: spacer ─────────────────────────────────────────────────
    ws.row_dimensions[2].height = 6

    # ── Row 3: Section header ─────────────────────────────────────────
    draw_section(ws, row=3, end_col=_END_COL, text="مكوّنات التقييم")

    # ── Row 4: Table header ───────────────────────────────────────────
    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        style_table_header(cell)
    ws.row_dimensions[4].height = 28

    # ── Data rows ─────────────────────────────────────────────────────
    data_start = 5
    for idx, av in enumerate(valuations):
        row = data_start + idx
        raw = (raw_components[idx] if raw_components and idx < len(raw_components) else {})
        area = raw.get("area_sqm")
        rate = raw.get("base_rate_per_sqm")

        row_values = [
            idx + 1,
            av.component_id,
            av.name,
            av.asset_type,
            av.purpose,
            area,
            rate,
            av.baseline_value,
            av.adjusted_value,
        ]
        alt_fill = get_fill(Palette.GREY_50) if idx % 2 else None
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if col in (8, 9):
                style_value(cell)
                cell.number_format = NumFormat.CURRENCY
            else:
                style_body(cell)
                if col == 6 and value is not None:
                    cell.number_format = NumFormat.QUANTITY
                elif col == 7 and value is not None:
                    cell.number_format = NumFormat.CURRENCY
            if alt_fill is not None:
                cell.fill = alt_fill
        ws.row_dimensions[row].height = 22

    # ── Totals row ────────────────────────────────────────────────────
    totals_row = data_start + len(valuations)
    total_baseline = sum(av.baseline_value for av in valuations)
    total_adjusted = sum(av.adjusted_value for av in valuations)

    _navy_font = get_font(bold=True, color=Palette.WHITE)
    _navy_fill = get_fill(Palette.NAVY)
    _gold_font = get_font(bold=True, color=Palette.GOLD_LIGHT)
    _center_al = get_alignment(h="center")

    for col in range(1, _END_COL + 1):
        c = ws.cell(row=totals_row, column=col)
        c.fill = _navy_fill
        c.alignment = _center_al
        if col == 1:
            c.value = "الإجمالي"
            c.font = _navy_font
        elif col == 8:
            c.value = total_baseline
            c.font = _gold_font
            c.number_format = NumFormat.CURRENCY
        elif col == 9:
            c.value = total_adjusted
            c.font = _gold_font
            c.number_format = NumFormat.CURRENCY
        else:
            c.font = _navy_font
    ws.row_dimensions[totals_row].height = 28

    # ── Validation section (optional) ─────────────────────────────────
    if validation is None or validation.is_clean:
        return

    next_row = totals_row + 2
    section_bg = Palette.CORAL if validation.is_blocking else Palette.INDIGO
    draw_section(ws, row=next_row, end_col=_END_COL, text="ملاحظات التحقق", bg=section_bg)
    next_row += 1

    v_headers = ["الخطورة", "الكود", "مكوّن #", "الحقل", "الرسالة"]
    for col, h in enumerate(v_headers, start=1):
        c = ws.cell(row=next_row, column=col, value=h)
        style_table_header(c)
    ws.row_dimensions[next_row].height = 24
    next_row += 1

    for issue in validation.issues:
        sev = issue.severity.value if hasattr(issue.severity, "value") else str(issue.severity)
        comp_idx = issue.component_index if issue.component_index is not None else "—"
        for col, val in enumerate(
            [sev, issue.code, comp_idx, issue.field, issue.message], start=1
        ):
            c = ws.cell(row=next_row, column=col, value=val)
            style_body(c)
        ws.row_dimensions[next_row].height = 20
        next_row += 1
