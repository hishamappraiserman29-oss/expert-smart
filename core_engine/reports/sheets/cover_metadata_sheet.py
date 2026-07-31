"""
cover_metadata_sheet.py — Cover & Metadata Excel sheet builder (Phase 14).

Renders firm identity, contact block, and report metadata onto a worksheet.
Display-only: no valuation math performed here.
"""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

try:
    from reports.report_identity import (
        APPRAISER_EMAIL,
        APPRAISER_NAME,
        APPRAISER_TEL,
        FIRM_NAME,
        ReportOutputMetadata,
    )
except ImportError:
    from report_identity import (  # type: ignore[no-redef]
        APPRAISER_EMAIL,
        APPRAISER_NAME,
        APPRAISER_TEL,
        FIRM_NAME,
        ReportOutputMetadata,
    )

# ── Style constants ───────────────────────────────────────────────────────────

_FONT_FIRM    = Font(bold=True, size=16)
_FONT_TITLE   = Font(bold=True, size=13)
_FONT_HEADER  = Font(bold=True, size=11)
_FONT_LABEL   = Font(bold=True, size=10)
_FONT_VALUE   = Font(size=10)
_FONT_CONTACT = Font(size=10, italic=True)
_FONT_BADGE   = Font(bold=True, size=9, color="FF0000")

_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_ALIGN_RIGHT  = Alignment(horizontal="right",  vertical="center", wrap_text=True)

_BORDER_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _set(ws: Worksheet, row: int, col: int, value: Any,
         font: Font | None = None,
         align: Alignment | None = None,
         border: Border | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if font   is not None: cell.font      = font
    if align  is not None: cell.alignment = align
    if border is not None: cell.border    = border


# ── Public API ────────────────────────────────────────────────────────────────

def apply_cover_metadata_sheet(
    ws: Worksheet,
    metadata: "ReportOutputMetadata | None" = None,
    *,
    extra: "dict[str, Any] | None" = None,
) -> None:
    """Write the Cover & Metadata sheet.

    Args:
        ws:       Target worksheet (caller sets the tab name).
        metadata: ReportOutputMetadata instance; None renders defaults.
        extra:    Optional dict of additional key/value rows to append.
    """
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 48

    row = 1

    # ── Firm identity banner ──────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    _set(ws, row, 1, FIRM_NAME, font=_FONT_FIRM, align=_ALIGN_CENTER)
    row += 1

    ws.row_dimensions[row].height = 16
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    _set(ws, row, 1, "تقرير التقييم العقاري", font=_FONT_TITLE, align=_ALIGN_CENTER)
    row += 1

    # ── Internal-only badge (when applicable) ─────────────────────────────────
    if metadata and metadata.excel_internal_only:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        _set(ws, row, 1,
             "⚠ INTERNAL USE ONLY — NOT FOR DISTRIBUTION TO EXTERNAL CLIENTS",
             font=_FONT_BADGE, align=_ALIGN_CENTER)
        row += 1

    row += 1  # spacer

    # ── Contact block ─────────────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    _set(ws, row, 1, "بيانات الخبير والتواصل", font=_FONT_HEADER, align=_ALIGN_CENTER)
    row += 1

    contact_rows = [
        ("خبير التقييم",  APPRAISER_NAME),
        ("Tel / WhatsApp", APPRAISER_TEL),
        ("Email",          APPRAISER_EMAIL),
    ]
    for label, value in contact_rows:
        _set(ws, row, 1, label, font=_FONT_LABEL,   align=_ALIGN_RIGHT, border=_BORDER_THIN)
        _set(ws, row, 2, value, font=_FONT_CONTACT, align=_ALIGN_LEFT,  border=_BORDER_THIN)
        row += 1

    row += 1  # spacer

    # ── Report metadata ───────────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    _set(ws, row, 1, "بيانات التقرير", font=_FONT_HEADER, align=_ALIGN_CENTER)
    row += 1

    if metadata is not None:
        meta_rows: list[tuple[str, Any]] = [
            ("report_id",            metadata.report_id),
            ("report_type",          metadata.report_type),
            ("output_audience",      metadata.output_audience),
            ("file_format",          metadata.file_format),
            ("generated_at",         metadata.generated_at),
            ("appraiser_name",       metadata.appraiser_name),
            ("valuation_purpose",    metadata.valuation_purpose),
            ("asset_type",           metadata.asset_type),
            ("asset_family",         metadata.asset_family),
            ("approval_status",      metadata.approval_status),
            ("source_log_available", str(metadata.source_log_available)),
            ("excel_internal_only",  str(metadata.excel_internal_only)),
        ]
    else:
        meta_rows = [
            ("report_id",            "—"),
            ("output_audience",      "—"),
            ("file_format",          "—"),
            ("generated_at",         "—"),
            ("appraiser_name",       APPRAISER_NAME),
            ("approval_status",      "not_required"),
            ("source_log_available", "False"),
            ("excel_internal_only",  "False"),
        ]

    for label, value in meta_rows:
        _set(ws, row, 1, label, font=_FONT_LABEL, align=_ALIGN_RIGHT, border=_BORDER_THIN)
        _set(ws, row, 2, str(value) if value is not None else "—",
             font=_FONT_VALUE, align=_ALIGN_LEFT, border=_BORDER_THIN)
        row += 1

    # ── Extra rows ────────────────────────────────────────────────────────────
    if extra:
        row += 1
        for label, value in extra.items():
            _set(ws, row, 1, str(label), font=_FONT_LABEL,  align=_ALIGN_RIGHT, border=_BORDER_THIN)
            _set(ws, row, 2, str(value), font=_FONT_VALUE,  align=_ALIGN_LEFT,  border=_BORDER_THIN)
            row += 1
