"""
formula_library_sheet.py — Formula Library Excel sheet builder (Phase 14).

DISPLAY ONLY — no valuation formulas are computed here.
Provides a structured reference sheet documenting the calculation methods
used in each valuation approach. Actual computation happens in the engines.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ── Style constants ───────────────────────────────────────────────────────────

_FONT_TITLE    = Font(bold=True, size=14)
_FONT_SECTION  = Font(bold=True, size=11, color="FFFFFF")
_FONT_LABEL    = Font(bold=True, size=10)
_FONT_FORMULA  = Font(size=10, name="Courier New")
_FONT_NOTE     = Font(size=9, italic=True, color="595959")

_FILL_SECTION  = PatternFill("solid", fgColor="1B3263")   # Navy — Midnight Gold palette
_FILL_METHOD   = PatternFill("solid", fgColor="294479")   # Navy light
_FILL_ALT_ROW  = PatternFill("solid", fgColor="EEF1F7")   # Grey 100

_ALIGN_CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT    = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

_BORDER_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


# ── Method reference data (display only) ─────────────────────────────────────

_METHODS: list[dict[str, str]] = [
    {
        "section":  "Sales Comparison Approach",
        "method":   "Adjusted Sales Price",
        "display":  "V = Σ(Comparable_Sale_Price × Adjustment_Factor) / n",
        "note":     "Adjustments for location, size, condition, time. Engine: adapters/residential, commercial.",
    },
    {
        "section":  "Sales Comparison Approach",
        "method":   "Price per Square Meter",
        "display":  "V = Adjusted_Price_per_m² × Subject_Area",
        "note":     "Used for standardised unit comparisons.",
    },
    {
        "section":  "Cost Approach",
        "method":   "Depreciated Replacement Cost (DRC)",
        "display":  "V = RCN − Accumulated_Depreciation + Land_Value",
        "note":     "RCN = Replacement Cost New. Depreciation: physical, functional, external.",
    },
    {
        "section":  "Cost Approach",
        "method":   "Insurance / Reinstatement",
        "display":  "V_ins = RCN + Professional_Fees + Demolition_Allowance",
        "note":     "Day-1 reinstatement basis. No land. Refer to insurance_adapter.",
    },
    {
        "section":  "Income Approach",
        "method":   "Direct Capitalisation",
        "display":  "V = NOI / Cap_Rate",
        "note":     "NOI = Gross Income − Vacancy − OpEx. Cap_Rate from market evidence.",
    },
    {
        "section":  "Income Approach",
        "method":   "Gross Rent Multiplier (GRM)",
        "display":  "V = Gross_Annual_Rent × GRM",
        "note":     "Quick cross-check only; not primary method.",
    },
    {
        "section":  "DCF / Investment Analysis",
        "method":   "Discounted Cash Flow (DCF)",
        "display":  "V = Σ [CF_t / (1+r)^t] + Terminal_Value / (1+r)^n",
        "note":     "r = discount rate (WACC/IRR target). Engine: adapters/dcf_model.",
    },
    {
        "section":  "DCF / Investment Analysis",
        "method":   "Internal Rate of Return (IRR)",
        "display":  "NPV = 0 ⟹ solve for r",
        "note":     "Sensitivity: adapters/dcf_sensitivity.",
    },
    {
        "section":  "HABU (Highest & Best Use)",
        "method":   "HABU Analysis",
        "display":  "Select use: max(V_use_1, V_use_2, …, V_use_n)",
        "note":     "Must be legal, physical, financially feasible, maximally productive.",
    },
    {
        "section":  "Residual Land Value",
        "method":   "Residual / Hypothetical Development",
        "display":  "V_land = GDV − Build_Cost − Developer_Profit − Finance_Costs",
        "note":     "GDV = Gross Development Value at completion.",
    },
    {
        "section":  "Mortgage / Lending Risk",
        "method":   "Loan-to-Value (LTV)",
        "display":  "LTV = Loan_Amount / Appraised_Value",
        "note":     "Basel III/IV thresholds apply. Refer to PURPOSE_RULES (protected).",
    },
    {
        "section":  "Mortgage / Lending Risk",
        "method":   "Mortgage Lending Value (MLV)",
        "display":  "MLV = MV × (1 − Haircut_%)",
        "note":     "Haircut per Basel IV; central bank overlay applied if present.",
    },
    {
        "section":  "Liquidation",
        "method":   "Forced Sale / Orderly Liquidation",
        "display":  "V_liq = MV × Liquidation_Factor (0.6–0.85 typical)",
        "note":     "Factor depends on asset type, liquidity, urgency.",
    },
    {
        "section":  "IFRS 13",
        "method":   "Fair Value Hierarchy",
        "display":  "Level 1: quoted prices | Level 2: observable inputs | Level 3: unobservable",
        "note":     "Engine: adapters/ifrs_13. Disclosure required.",
    },
    {
        "section":  "Property Rights / Usufruct / Minority",
        "method":   "Leasehold / Usufruct Discount",
        "display":  "V_leasehold = V_freehold × (1 − Yield_Capitalisation_Factor)",
        "note":     "Duration-weighted. Minority interest discount applied separately.",
    },
    {
        "section":  "ESG / Remediation",
        "method":   "Environmental Adjustment",
        "display":  "V_adj = V_unimpaired − Remediation_Cost − Stigma_Discount",
        "note":     "Stigma may persist post-remediation.",
    },
    {
        "section":  "Litigation / Compensation",
        "method":   "Expropriation / Compulsory Purchase",
        "display":  "Compensation = MV + Disturbance + Loss_of_Profit + Severance",
        "note":     "Statutory basis; jurisdiction-specific rules apply.",
    },
    {
        "section":  "Reconciliation",
        "method":   "Weighted Reconciliation",
        "display":  "V_final = Σ(Method_Value_i × Weight_i)  where Σ Weight_i = 1",
        "note":     "Weights set by appraiser judgement. Engine: reports/sheets/reconciliation.",
    },
]


# ── Public API ────────────────────────────────────────────────────────────────

def apply_formula_library_sheet(ws: Worksheet) -> None:
    """Write the Formula Library reference sheet.

    Args:
        ws: Target worksheet (caller names the tab).

    DISPLAY ONLY — no calculations performed.
    """
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 50

    row = 1

    # Title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="Formula Library — Reference Only (Display)")
    cell.font      = _FONT_TITLE
    cell.alignment = _ALIGN_CENTER
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note = ws.cell(row=row, column=1,
                   value="⚠ Formulas below are DESCRIPTIVE ONLY. Actual computation is performed by the valuation engines.")
    note.font      = _FONT_NOTE
    note.alignment = _ALIGN_LEFT
    row += 1
    row += 1  # spacer

    # Column headers
    headers = ["Section / Approach", "Method", "Formula (Display)", "Notes / Engine Reference"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font      = _FONT_SECTION
        cell.fill      = _FILL_SECTION
        cell.alignment = _ALIGN_CENTER
        cell.border    = _BORDER_THIN
    row += 1

    prev_section = ""
    alt = False

    for entry in _METHODS:
        section = entry["section"]
        if section != prev_section:
            # Section sub-header row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            sh = ws.cell(row=row, column=1, value=section)
            sh.font      = Font(bold=True, size=10, color="FFFFFF")
            sh.fill      = _FILL_METHOD
            sh.alignment = _ALIGN_LEFT
            sh.border    = _BORDER_THIN
            row += 1
            prev_section = section
            alt = False

        fill = _FILL_ALT_ROW if alt else None
        values = ["", entry["method"], entry["display"], entry["note"]]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font      = _FONT_FORMULA if col == 3 else _FONT_NOTE if col == 4 else _FONT_LABEL
            cell.alignment = _ALIGN_LEFT
            cell.border    = _BORDER_THIN
            if fill:
                cell.fill = fill
        alt = not alt
        row += 1
