"""
human_approval_log_sheet.py — Human Approval Log Excel sheet builder (Phase 14).

Records the human approval state for auto-enriched fields.
Linked to Phase 13 approval_rules: APPROVAL_STATUSES, approval_status,
requires_human_approval, get_report_status_from_approval.
Audit/governance layer only — no valuation logic.
"""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ── Style constants ───────────────────────────────────────────────────────────

_FONT_TITLE     = Font(bold=True, size=14)
_FONT_HEADER    = Font(bold=True, size=10, color="FFFFFF")
_FONT_LABEL     = Font(bold=True, size=10)
_FONT_VALUE     = Font(size=10)
_FONT_NOTE      = Font(size=9, italic=True, color="595959")
_FONT_APPROVED  = Font(size=10, bold=True, color="0E8B6E")
_FONT_PENDING   = Font(size=10, bold=True, color="C9A961")
_FONT_REJECTED  = Font(size=10, bold=True, color="D85842")

_FILL_HEADER    = PatternFill("solid", fgColor="6B4FA1")   # Purple (governance)
_FILL_APPROVED  = PatternFill("solid", fgColor="D4EDDA")
_FILL_PENDING   = PatternFill("solid", fgColor="FFF3CD")
_FILL_REJECTED  = PatternFill("solid", fgColor="F8D7DA")
_FILL_ALT       = PatternFill("solid", fgColor="EEF1F7")

_ALIGN_CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT     = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

_BORDER_THIN    = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

_COLUMNS: list[tuple[str, str, int]] = [
    ("field_code",        "Field",                22),
    ("is_automated_fill", "Auto-Enriched?",       16),
    ("confidence_score",  "Confidence (%)",       16),
    ("approval_status",   "Approval Status",      22),
    ("reviewer",          "Reviewer",             22),
    ("reviewed_at",       "Reviewed At",          20),
    ("notes",             "Notes",                36),
]

_STATUS_STYLES: dict[str, tuple[Font, PatternFill]] = {
    "approved":              (_FONT_APPROVED, _FILL_APPROVED),
    "pending_human_review":  (_FONT_PENDING,  _FILL_PENDING),
    "rejected":              (_FONT_REJECTED, _FILL_REJECTED),
    "not_required":          (_FONT_VALUE,    _FILL_ALT),
}


def apply_human_approval_log_sheet(
    ws: Worksheet,
    approval_data: "dict[str, Any] | None" = None,
    entries: "list[dict[str, Any]] | None" = None,
) -> None:
    """Write the Human Approval Log sheet.

    Args:
        ws:            Target worksheet.
        approval_data: Top-level approval data dict (from Phase 13 fields).
        entries:       Optional list of per-field approval records.
                       Each dict may have keys matching _COLUMNS field codes.
    """
    for col_idx, (_, _, width) in enumerate(_COLUMNS, 1):
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = width

    row = 1

    # Title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
    cell = ws.cell(row=row, column=1, value="Human Approval Log — Auto-Enrichment Governance")
    cell.font      = _FONT_TITLE
    cell.alignment = _ALIGN_CENTER
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
    note = ws.cell(row=row, column=1,
                   value="Linked to Phase 13 approval_rules. Auto-enriched values require "
                         "approval_status='approved' before entering a final report.")
    note.font      = _FONT_NOTE
    note.alignment = _ALIGN_LEFT
    row += 1

    # Overall approval status summary
    if approval_data:
        row += 1
        summary_items = [
            ("is_automated_fill",  approval_data.get("is_automated_fill", "—")),
            ("approval_status",    approval_data.get("approval_status",   "—")),
            ("confidence_score",   approval_data.get("confidence_score",  "—")),
            ("source_log",         str(bool(approval_data.get("data_source_log")))),
        ]
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
        sh = ws.cell(row=row, column=1, value="Overall Approval Status")
        sh.font      = _FONT_LABEL
        sh.alignment = _ALIGN_LEFT
        row += 1

        for label, value in summary_items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            lc = ws.cell(row=row, column=1, value=label)
            lc.font = _FONT_LABEL
            lc.border = _BORDER_THIN
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=len(_COLUMNS))
            vc = ws.cell(row=row, column=3, value=str(value))
            vc.font = _FONT_VALUE
            vc.border = _BORDER_THIN
            row += 1

    row += 1  # spacer

    # Column headers
    for col_idx, (_, label, _w) in enumerate(_COLUMNS, 1):
        cell = ws.cell(row=row, column=col_idx, value=label)
        cell.font      = _FONT_HEADER
        cell.fill      = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border    = _BORDER_THIN
    row += 1

    # Per-field entries
    field_entries = entries or []
    for entry in field_entries:
        status = str(entry.get("approval_status", "not_required"))
        style_font, style_fill = _STATUS_STYLES.get(status, (_FONT_VALUE, _FILL_ALT))

        for col_idx, (field_code, _, _w) in enumerate(_COLUMNS, 1):
            value = entry.get(field_code, "")
            cell = ws.cell(row=row, column=col_idx, value=str(value) if value else "")
            cell.font      = style_font if field_code == "approval_status" else _FONT_VALUE
            cell.fill      = style_fill if field_code == "approval_status" else None
            cell.alignment = _ALIGN_LEFT
            cell.border    = _BORDER_THIN
        row += 1

    if not field_entries:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
        ph = ws.cell(row=row, column=1,
                     value="[No per-field approval entries — populate from agentic enrichment results]")
        ph.font      = _FONT_NOTE
        ph.alignment = _ALIGN_LEFT
