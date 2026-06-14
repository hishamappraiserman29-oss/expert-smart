"""
source_log_sheet.py — Source Log Excel sheet builder (Phase 14).

Provides a structured sheet for recording data sources used in the valuation.
Links to Phase 13 source_log_available / data_source_log metadata fields.
Display/audit layer only — no valuation logic.
"""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ── Style constants ───────────────────────────────────────────────────────────

_FONT_TITLE   = Font(bold=True, size=14)
_FONT_HEADER  = Font(bold=True, size=10, color="FFFFFF")
_FONT_LABEL   = Font(bold=True, size=10)
_FONT_VALUE   = Font(size=10)
_FONT_NOTE    = Font(size=9, italic=True, color="595959")

_FILL_HEADER  = PatternFill("solid", fgColor="0E8B6E")   # Emerald
_FILL_ALT     = PatternFill("solid", fgColor="EEF1F7")   # Grey 100

_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

_BORDER_THIN  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

_COLUMNS: list[tuple[str, str, int]] = [
    ("source_id",      "Source ID",            14),
    ("source_type",    "Source Type",          22),
    ("source_name",    "Source Name / URL",    40),
    ("retrieval_date", "Retrieval Date",        18),
    ("field_applied",  "Field Applied To",      24),
    ("confidence",     "Confidence (%)",        16),
    ("notes",          "Notes",                 36),
]


def apply_source_log_sheet(
    ws: Worksheet,
    source_log: "list[dict[str, Any]] | None" = None,
) -> None:
    """Write the Source Log sheet.

    Args:
        ws:         Target worksheet.
        source_log: Optional list of source dicts from Phase 13 data_source_log.
                    Each dict may have keys matching _COLUMNS field codes.
    """
    for col_idx, (_, label, width) in enumerate(_COLUMNS, 1):
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = width

    row = 1

    # Title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
    cell = ws.cell(row=row, column=1, value="Source Log — Data Provenance Audit Trail")
    cell.font      = _FONT_TITLE
    cell.alignment = _ALIGN_CENTER
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
    note = ws.cell(row=row, column=1,
                   value="Records all external and automated data sources used in this valuation. "
                         "Linked to Phase 13 source_log_available / data_source_log metadata.")
    note.font      = _FONT_NOTE
    note.alignment = _ALIGN_LEFT
    row += 1
    row += 1  # spacer

    # Column headers
    for col_idx, (_, label, _w) in enumerate(_COLUMNS, 1):
        cell = ws.cell(row=row, column=col_idx, value=label)
        cell.font      = _FONT_HEADER
        cell.fill      = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border    = _BORDER_THIN
    row += 1

    # Data rows
    entries = source_log or []
    for i, entry in enumerate(entries):
        fill = _FILL_ALT if i % 2 == 1 else None
        for col_idx, (field_code, _, _w) in enumerate(_COLUMNS, 1):
            value = entry.get(field_code, "")
            cell = ws.cell(row=row, column=col_idx, value=str(value) if value else "")
            cell.font      = _FONT_VALUE
            cell.alignment = _ALIGN_LEFT
            cell.border    = _BORDER_THIN
            if fill:
                cell.fill = fill
        row += 1

    # If no entries, show placeholder row
    if not entries:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(_COLUMNS))
        ph = ws.cell(row=row, column=1, value="[No source log entries — populate from data_source_log field]")
        ph.font      = _FONT_NOTE
        ph.alignment = _ALIGN_LEFT
        row += 1
