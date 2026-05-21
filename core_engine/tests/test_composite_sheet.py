"""
Tests for core_engine/reports/sheets/composite_sheet.py (Wave 6).

Coverage:
  1. Sheet is created with the correct name.
  2. Banner row contains expected text.
  3. Table header row contains all required column headers.
  4. Data rows: component fields written correctly.
  5. Totals row: sum of baseline and adjusted values.
  6. Validation section present when issues exist, absent when clean.
  7. Regression: ExcelReportBuilder.build() without composite data is unchanged.
  8. ExcelReportBuilder.build() creates the sheet only when composite data
     is explicitly passed.
"""
from __future__ import annotations

import os
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest

# ── sys.path ──────────────────────────────────────────────────────────
_CORE = Path(__file__).resolve().parents[1]
_ROOT = _CORE.parent

for _p in (str(_CORE), str(_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

os.chdir(str(_CORE))

# ── imports under test ────────────────────────────────────────────────
from openpyxl import Workbook

from reports.sheets.composite_sheet import build_composite_sheet, _SHEET_NAME
from adapters.purpose_adapter import AdjustedValuation
from validation.composite_rules import (
    CompositeValidationReport,
    Severity,
    ValidationIssue,
)


# ─────────────────────────────────────────────────────────────────────
# Helpers / fixtures
# ─────────────────────────────────────────────────────────────────────

def _make_av(
    idx: int = 0,
    baseline: float = 1_000_000.0,
    adjusted: float = 950_000.0,
    purpose: str = "البيع والشراء - القيمة السوقية العادلة (Market Value)",
    asset_type: str = "وحدة سكنية (شقة / فيلا)",
    deep: bool = False,
    iaao: bool = False,
) -> AdjustedValuation:
    return AdjustedValuation(
        component_id=f"c{idx}",
        name=f"وحدة {idx + 1}",
        asset_type=asset_type,
        purpose=purpose,
        baseline_value=baseline,
        multiplier_applied=round(adjusted / baseline, 4) if baseline else 1.0,
        adjusted_value=adjusted,
        route="market_baseline",
        deep_route_deferred=deep,
        iaao_block_triggered=iaao,
        uspap_standards=("Standard 1", "Standard 2"),
        notes=(),
    )


def _make_raw(idx: int = 0, area: float = 100.0, rate: float = 10_000.0) -> dict:
    return {
        "id": f"c{idx}",
        "area_sqm": area,
        "base_rate_per_sqm": rate,
    }


def _clean_report() -> CompositeValidationReport:
    return CompositeValidationReport(issues=())


def _report_with_advisory() -> CompositeValidationReport:
    return CompositeValidationReport(issues=(
        ValidationIssue(
            severity=Severity.ADVISORY,
            code="ADV001",
            component_index=0,
            field="area_sqm",
            message="مساحة صغيرة جداً",
        ),
    ))


def _report_with_blocking() -> CompositeValidationReport:
    return CompositeValidationReport(issues=(
        ValidationIssue(
            severity=Severity.BLOCKING,
            code="BLK001",
            component_index=0,
            field="asset_type",
            message="نوع أصل غير معروف",
        ),
    ))


# ─────────────────────────────────────────────────────────────────────
# 1. Sheet creation
# ─────────────────────────────────────────────────────────────────────

class TestSheetCreation:
    def test_sheet_name_correct(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        assert _SHEET_NAME in wb.sheetnames

    def test_sheet_is_last_tab(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        assert wb.sheetnames[-1] == _SHEET_NAME

    def test_rtl_view_enabled(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        ws = wb[_SHEET_NAME]
        assert ws.sheet_view.rightToLeft is True


# ─────────────────────────────────────────────────────────────────────
# 2. Banner
# ─────────────────────────────────────────────────────────────────────

class TestBanner:
    def test_banner_row1_contains_title(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        ws = wb[_SHEET_NAME]
        assert ws["A1"].value is not None
        assert "التقييم المركب" in str(ws["A1"].value)


# ─────────────────────────────────────────────────────────────────────
# 3. Table header row
# ─────────────────────────────────────────────────────────────────────

class TestTableHeader:
    def _header_values(self, wb):
        ws = wb[_SHEET_NAME]
        return [ws.cell(row=4, column=c).value for c in range(1, 10)]

    def test_serial_column_header(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        assert self._header_values(wb)[0] == "م"

    def test_baseline_column_header_present(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        headers = self._header_values(wb)
        assert any("أساسية" in str(h) for h in headers)

    def test_adjusted_column_header_present(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        headers = self._header_values(wb)
        assert any("معدّلة" in str(h) for h in headers)

    def test_nine_header_columns(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        headers = self._header_values(wb)
        assert all(h is not None for h in headers)


# ─────────────────────────────────────────────────────────────────────
# 4. Data rows
# ─────────────────────────────────────────────────────────────────────

class TestDataRows:
    def test_single_component_row_written(self):
        wb = Workbook()
        av = _make_av(0, baseline=2_000_000.0, adjusted=1_900_000.0)
        build_composite_sheet(wb, [av])
        ws = wb[_SHEET_NAME]
        # Row 4 = headers, row 5 = first data row
        assert ws.cell(row=5, column=1).value == 1       # serial
        assert ws.cell(row=5, column=2).value == "c0"    # component_id
        assert ws.cell(row=5, column=8).value == pytest.approx(2_000_000.0)
        assert ws.cell(row=5, column=9).value == pytest.approx(1_900_000.0)

    def test_two_components_written(self):
        wb = Workbook()
        avs = [_make_av(0), _make_av(1, baseline=500_000, adjusted=475_000)]
        build_composite_sheet(wb, avs)
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=1).value == 1
        assert ws.cell(row=6, column=1).value == 2

    def test_asset_type_in_column_4(self):
        wb = Workbook()
        av = _make_av(0, asset_type="أرض سكنية")
        build_composite_sheet(wb, [av])
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=4).value == "أرض سكنية"

    def test_purpose_in_column_5(self):
        wb = Workbook()
        purpose = "البيع والشراء - القيمة السوقية العادلة (Market Value)"
        av = _make_av(0, purpose=purpose)
        build_composite_sheet(wb, [av])
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=5).value == purpose

    def test_area_sqm_from_raw_components(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], raw_components=[_make_raw(area=250.0)])
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=6).value == pytest.approx(250.0)

    def test_base_rate_from_raw_components(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], raw_components=[_make_raw(rate=8_500.0)])
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=7).value == pytest.approx(8_500.0)

    def test_area_rate_none_when_no_raw_components(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        ws = wb[_SHEET_NAME]
        assert ws.cell(row=5, column=6).value is None
        assert ws.cell(row=5, column=7).value is None


# ─────────────────────────────────────────────────────────────────────
# 5. Totals row
# ─────────────────────────────────────────────────────────────────────

class TestTotalsRow:
    def test_totals_label(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()])
        ws = wb[_SHEET_NAME]
        totals_row = 5 + 1   # data_start + 1 component
        assert ws.cell(row=totals_row, column=1).value == "الإجمالي"

    def test_total_baseline_sum(self):
        wb = Workbook()
        avs = [_make_av(0, baseline=1_000_000), _make_av(1, baseline=500_000)]
        build_composite_sheet(wb, avs)
        ws = wb[_SHEET_NAME]
        totals_row = 5 + 2
        assert ws.cell(row=totals_row, column=8).value == pytest.approx(1_500_000)

    def test_total_adjusted_sum(self):
        wb = Workbook()
        avs = [
            _make_av(0, baseline=1_000_000, adjusted=950_000),
            _make_av(1, baseline=500_000,   adjusted=410_000),
        ]
        build_composite_sheet(wb, avs)
        ws = wb[_SHEET_NAME]
        totals_row = 5 + 2
        assert ws.cell(row=totals_row, column=9).value == pytest.approx(1_360_000)

    def test_empty_valuations_totals_zero(self):
        wb = Workbook()
        build_composite_sheet(wb, [])
        ws = wb[_SHEET_NAME]
        totals_row = 5
        assert ws.cell(row=totals_row, column=8).value == pytest.approx(0.0)
        assert ws.cell(row=totals_row, column=9).value == pytest.approx(0.0)


# ─────────────────────────────────────────────────────────────────────
# 6. Validation section
# ─────────────────────────────────────────────────────────────────────

class TestValidationSection:
    def _last_section_text(self, wb):
        ws = wb[_SHEET_NAME]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and "ملاحظات التحقق" in str(cell.value):
                    return cell.value
        return None

    def test_no_validation_section_when_clean(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=_clean_report())
        assert self._last_section_text(wb) is None

    def test_no_validation_section_when_none(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=None)
        assert self._last_section_text(wb) is None

    def test_advisory_section_appears(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=_report_with_advisory())
        assert self._last_section_text(wb) is not None

    def test_blocking_section_appears(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=_report_with_blocking())
        assert self._last_section_text(wb) is not None

    def test_issue_code_written(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=_report_with_advisory())
        ws = wb[_SHEET_NAME]
        codes = [c.value for row in ws.iter_rows() for c in row if c.value == "ADV001"]
        assert codes, "Issue code 'ADV001' not found in sheet"

    def test_issue_message_written(self):
        wb = Workbook()
        build_composite_sheet(wb, [_make_av()], validation=_report_with_advisory())
        ws = wb[_SHEET_NAME]
        msgs = [c.value for row in ws.iter_rows() for c in row if c.value == "مساحة صغيرة جداً"]
        assert msgs, "Issue message not found in sheet"


# ─────────────────────────────────────────────────────────────────────
# 7. ExcelReportBuilder regression — no composite data
# ExcelReportBuilder removes the default blank sheet in __init__,
# so result=None leaves an empty workbook.  We add a sentinel sheet
# before calling build() so openpyxl can save without error.
# ─────────────────────────────────────────────────────────────────────

class TestExcelBuilderRegression:
    @pytest.fixture(autouse=True)
    def _env(self, monkeypatch, tmp_path):
        monkeypatch.setenv("JWT_SECRET", "test-wave6-builder")
        self.tmp = tmp_path

    def _minimal_builder(self):
        from reports.excel_builder import ExcelReportBuilder
        builder = ExcelReportBuilder(result=None)
        builder.workbook.create_sheet("_sentinel")   # keeps workbook saveable
        return builder

    def test_build_without_composite_runs_without_error(self):
        builder = self._minimal_builder()
        out = builder.build(str(self.tmp / "test_no_composite.xlsx"))
        assert Path(out).exists()

    def test_composite_sheet_absent_without_composite_data(self):
        builder = self._minimal_builder()
        out = builder.build(str(self.tmp / "test_no_composite2.xlsx"))
        import openpyxl
        wb = openpyxl.load_workbook(out)
        assert "التقييم المركب" not in wb.sheetnames


# ─────────────────────────────────────────────────────────────────────
# 8. ExcelReportBuilder — sheet added only with composite data
# ─────────────────────────────────────────────────────────────────────

class TestExcelBuilderCompositeHook:
    @pytest.fixture(autouse=True)
    def _env(self, monkeypatch, tmp_path):
        monkeypatch.setenv("JWT_SECRET", "test-wave6-hook")
        self.tmp = tmp_path

    def _minimal_builder(self):
        from reports.excel_builder import ExcelReportBuilder
        builder = ExcelReportBuilder(result=None)
        builder.workbook.create_sheet("_sentinel")
        return builder

    def test_composite_sheet_present_when_data_passed(self):
        import openpyxl
        builder = self._minimal_builder()
        avs = [_make_av(0, baseline=1_000_000, adjusted=950_000)]
        out = builder.build(
            str(self.tmp / "test_with_composite.xlsx"),
            composite_valuations=avs,
        )
        wb = openpyxl.load_workbook(out)
        assert "التقييم المركب" in wb.sheetnames

    def test_composite_sheet_totals_correct_via_builder(self):
        import openpyxl
        builder = self._minimal_builder()
        avs = [
            _make_av(0, baseline=2_000_000, adjusted=1_900_000),
            _make_av(1, baseline=1_000_000, adjusted=820_000),
        ]
        out = builder.build(
            str(self.tmp / "test_composite_totals.xlsx"),
            composite_valuations=avs,
        )
        wb = openpyxl.load_workbook(out)
        ws = wb["التقييم المركب"]
        totals_row = 5 + 2
        assert ws.cell(row=totals_row, column=8).value == pytest.approx(3_000_000)
        assert ws.cell(row=totals_row, column=9).value == pytest.approx(2_720_000)
