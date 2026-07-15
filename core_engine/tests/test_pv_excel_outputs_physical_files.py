# test_pv_excel_outputs_physical_files.py
# Part H — Physical file verification for excel_outputs/ folder
# advisory_only=True | not_real_training=True | no_commit=True

import json
import pathlib
import pytest

QA     = pathlib.Path(
    "core_engine/instance/manual_review_outputs/"
    "professional_valuation_full_page_uat_legacy_excel_visual_review"
)
EXCEL_OUT   = QA / "excel_outputs"
AUDIT_JSON  = EXCEL_OUT / "00_excel_outputs_filesystem_audit.json"
FINAL_RPT   = QA / "final_report" / "final_professional_valuation_full_page_uat_report.txt"
LEGACY_REF  = QA / "legacy_references"
GAP_JSON    = QA / "excel_comparison_audits" / "generated_vs_legacy_excel_quality_gap_analysis.json"

GENERATED_FILES = [
    "traditional_report_workbook.xlsx",
    "detailed_report_workbook.xlsx",
    "professional_report_workbook.xlsx",
    "hbu_analysis_report_workbook.xlsx",
    "standards_compliance_report_workbook.xlsx",
    "report_review_output_workbook.xlsx",
    "simulated_uploaded_report_workbook.xlsx",
]

LEGACY_FILES = [
    "Report_ES_GRAND_FINAL_v4.xlsm",
    "Report_ES_ULTRA.xlsm",
]


# ── T01: QA folder exists ─────────────────────────────────────────────────────
def test_T01_qa_folder_exists():
    assert QA.exists(), f"QA folder missing: {QA}"


# ── T02: excel_outputs folder exists ─────────────────────────────────────────
def test_T02_excel_outputs_folder_exists():
    assert EXCEL_OUT.exists(), f"excel_outputs folder missing: {EXCEL_OUT}"


# ── T03: filesystem audit file exists ────────────────────────────────────────
def test_T03_filesystem_audit_exists():
    assert AUDIT_JSON.exists(), f"Filesystem audit missing: {AUDIT_JSON}"


# ── T04: filesystem audit has correct structure ───────────────────────────────
def test_T04_filesystem_audit_correct_structure():
    data = json.loads(AUDIT_JSON.read_text(encoding="utf-8"))
    assert data["excel_outputs_folder_exists"] is True
    assert data["excel_outputs_file_count"] == 7
    assert len(data["excel_outputs_files"]) == 7
    assert data["legacy_reference_files_are_NOT_counted_as_generated"] is True
    assert data["status"] == "GENERATED"


# ── T05: all 7 generated Excel files physically exist ────────────────────────
def test_T05_all_generated_excel_files_exist():
    for fname in GENERATED_FILES:
        f = EXCEL_OUT / fname
        assert f.exists(), f"Generated Excel file missing: {f}"
        assert f.stat().st_size > 0, f"Generated Excel file is empty: {f}"


# ── T06: generated files are readable by openpyxl ────────────────────────────
def test_T06_generated_files_readable_by_openpyxl():
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    for fname in GENERATED_FILES:
        f = EXCEL_OUT / fname
        if not f.exists():
            pytest.fail(f"File missing: {f}")
        wb = openpyxl.load_workbook(str(f), read_only=True)
        sheet_count = len(wb.sheetnames)
        wb.close()
        assert sheet_count >= 1, f"{fname} has no worksheets"


# ── T07: generated files have at least one sheet and meaningful content ───────
def test_T07_generated_files_have_content():
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    for fname in GENERATED_FILES:
        f = EXCEL_OUT / fname
        if not f.exists():
            pytest.fail(f"File missing: {f}")
        wb = openpyxl.load_workbook(str(f), read_only=True)
        # First sheet should have content
        ws = wb.active
        found_content = False
        for row in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                found_content = True
                break
        wb.close()
        assert found_content, f"{fname}: first sheet has no non-empty cells"


# ── T08: legacy reference files are NOT in excel_outputs ─────────────────────
def test_T08_legacy_files_not_in_excel_outputs():
    for fname in LEGACY_FILES:
        legacy_in_excel_out = EXCEL_OUT / fname
        assert not legacy_in_excel_out.exists(), (
            f"Legacy reference file should NOT be in excel_outputs/: {fname}"
        )


# ── T09: no generated file claims final certification ────────────────────────
def test_T09_no_fake_certification_in_generated_files():
    fake_phrases = ["certified_final_valuation", "FINAL_CERTIFIED", "شهادة معتمدة نهائية"]
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    for fname in GENERATED_FILES:
        f = EXCEL_OUT / fname
        if not f.exists():
            continue
        wb = openpyxl.load_workbook(str(f), read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is not None:
                        cell_str = str(v)
                        for phrase in fake_phrases:
                            assert phrase not in cell_str, (
                                f"Fake certification phrase '{phrase}' found in {fname}"
                            )
        wb.close()


# ── T10: no internal paths in generated files ─────────────────────────────────
def test_T10_no_internal_paths_in_generated_files():
    forbidden = [r"C:\Users\\", r"C:/Users/", "/home/", "AppData", "__file__"]
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl not installed")

    for fname in GENERATED_FILES:
        f = EXCEL_OUT / fname
        if not f.exists():
            continue
        wb = openpyxl.load_workbook(str(f), read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is not None:
                        cell_str = str(v)
                        for pat in forbidden:
                            assert pat not in cell_str, (
                                f"Internal path '{pat}' found in {fname}"
                            )
        wb.close()


# ── T11: final report includes Excel physical file verification section ───────
def test_T11_final_report_has_physical_verification_section():
    assert FINAL_RPT.exists(), f"Final report missing: {FINAL_RPT}"
    text = FINAL_RPT.read_text(encoding="utf-8")
    assert "Excel Outputs Physical File Verification" in text or \
           "PHYSICAL FILE VERIFICATION" in text or \
           "excel_outputs_filesystem_audit" in text, \
        "Final report must include Excel physical file verification section"


# ── T12: if report claims generated Excel, files must physically exist ────────
def test_T12_report_claims_match_physical_files():
    text = FINAL_RPT.read_text(encoding="utf-8")
    # If report says "GENERATED", at least one generated file must exist
    if "GENERATED" in text and "excel_outputs" in text:
        existing = [f for f in GENERATED_FILES if (EXCEL_OUT / f).exists()]
        assert len(existing) > 0, (
            "Final report claims generated Excel outputs but no files found in excel_outputs/"
        )


# ── T13: gap analysis is not marked COMPLETED without generated files ─────────
def test_T13_gap_analysis_completion_requires_physical_files():
    assert GAP_JSON.exists(), f"Gap analysis missing: {GAP_JSON}"
    data = json.loads(GAP_JSON.read_text(encoding="utf-8"))
    status = data.get("comparison_status", "")
    found = data.get("generated_excel_files_physically_found", False)
    if status == "COMPLETED":
        assert found is True, (
            "gap analysis says COMPLETED but generated_excel_files_physically_found is False"
        )
    if status == "BLOCKED":
        # No files should exist either
        existing = [f for f in GENERATED_FILES if (EXCEL_OUT / f).exists()]
        assert len(existing) == 0, (
            "gap analysis says BLOCKED but generated files exist in excel_outputs/"
        )


# ── T14: gap analysis documents critical gaps ─────────────────────────────────
def test_T14_gap_analysis_documents_critical_gaps():
    data = json.loads(GAP_JSON.read_text(encoding="utf-8"))
    assert data["legacy_references_used"] is True
    assert len(data.get("critical_gaps", [])) > 0, "Gap analysis must document critical gaps"
    assert len(data.get("recommended_next_fixes", [])) > 0


# ── T15: legacy ref files are present and sizable ────────────────────────────
def test_T15_legacy_reference_files_present():
    for fname in LEGACY_FILES:
        f = LEGACY_REF / fname
        assert f.exists(), f"Legacy reference missing: {f}"
        assert f.stat().st_size > 100_000, f"Legacy reference suspiciously small: {f}"
