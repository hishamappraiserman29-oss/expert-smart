# -*- coding: utf-8 -*-
"""
Visual QA tests — Excel workbook (73 sheets) + three-tier PDF system.
17 tests covering artifacts, governance, and cross-consistency.
"""
import json
import pathlib
import re

import pytest
import openpyxl

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
_BASE = pathlib.Path(__file__).resolve().parent.parent
_EXCEL_DIR = _BASE / "instance" / "manual_review_outputs" / "professional_valuation_excel_reference_parity" / "excel_outputs"
_EXCEL = _EXCEL_DIR / "core_valuation_master_workbook_reference_parity.xlsx"

_PDF_DIR = _BASE / "instance" / "manual_review_outputs" / "professional_valuation_pdf_reference_parity" / "pdf_outputs"
_TRAD_PDF = _PDF_DIR / "traditional_three_tier_report.pdf"
_DET_PDF  = _PDF_DIR / "detailed_three_tier_report.pdf"
_PRO_PDF  = _PDF_DIR / "professional_three_tier_report.pdf"

_TPL_DIR = _BASE / "templates" / "pdf"
_TRAD_TPL = _TPL_DIR / "pv_traditional_report.html"
_DET_TPL  = _TPL_DIR / "pv_detailed_report.html"
_PRO_TPL  = _TPL_DIR / "pv_professional_report.html"
_BLDR     = _BASE / "reports" / "pv_three_tier_pdf_builder.py"

_QA = _BASE / "instance" / "manual_review_outputs" / "professional_valuation_full_visual_qa"
_AUD = _QA / "audits"
_VI  = _QA / "visual_index"
_FR  = _QA / "final_report"

_EXCEL_QA_AUDIT  = _AUD / "excel_visual_qa_audit.json"
_PDF_PHYS_AUDIT  = _AUD / "pdf_physical_validation_audit.json"
_PDF_CONT_AUDIT  = _AUD / "pdf_content_validation_audit.json"
_CROSS_AUDIT     = _AUD / "excel_pdf_cross_consistency_audit.json"
_VI_INDEX        = _VI  / "OPEN_EXCEL_AND_PDF_VISUAL_QA.html"
_SUMMARY_JSON    = _FR  / "final_excel_pdf_visual_qa_summary.json"

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _pdf_text(path: pathlib.Path) -> str:
    try:
        import fitz
        doc = fitz.open(str(path))
        return "\n".join(doc[i].get_text() for i in range(doc.page_count))
    except Exception:
        return ""


def _read(path: pathlib.Path) -> str:
    if path.exists():
        try:
            return path.read_text(encoding="utf-8")
        except Exception:
            return ""
    return ""


def _triple(pdf: pathlib.Path, tpl: pathlib.Path) -> str:
    prv_dir = _BASE / "instance" / "manual_review_outputs" / "professional_valuation_pdf_reference_parity" / "visual_previews"
    tier = pdf.stem.split("_")[0]
    prv = prv_dir / f"{tier}_three_tier_preview.html"
    return _pdf_text(pdf) + "\n" + _read(prv) + "\n" + _read(tpl)


# ---------------------------------------------------------------------------
# 1. Excel workbook exists
# ---------------------------------------------------------------------------
def test_01_excel_workbook_exists():
    assert _EXCEL.exists(), str(_EXCEL)


# ---------------------------------------------------------------------------
# 2. Excel workbook opens
# ---------------------------------------------------------------------------
def test_02_excel_workbook_opens():
    wb = openpyxl.load_workbook(str(_EXCEL), read_only=True, data_only=True)
    assert wb.sheetnames
    wb.close()


# ---------------------------------------------------------------------------
# 3. Excel sheet count is 73
# ---------------------------------------------------------------------------
def test_03_excel_sheet_count_73():
    wb = openpyxl.load_workbook(str(_EXCEL), read_only=True, data_only=True)
    assert len(wb.sheetnames) == 73, f"expected 73, got {len(wb.sheetnames)}"
    wb.close()


# ---------------------------------------------------------------------------
# 4. Excel visual audit exists
# ---------------------------------------------------------------------------
def test_04_excel_visual_audit_exists():
    assert _EXCEL_QA_AUDIT.exists(), str(_EXCEL_QA_AUDIT)
    assert _EXCEL_QA_AUDIT.stat().st_size > 512


# ---------------------------------------------------------------------------
# 5. PDF physical audit exists
# ---------------------------------------------------------------------------
def test_05_pdf_physical_audit_exists():
    assert _PDF_PHYS_AUDIT.exists(), str(_PDF_PHYS_AUDIT)
    assert _PDF_PHYS_AUDIT.stat().st_size > 512


# ---------------------------------------------------------------------------
# 6. PDF content audit exists
# ---------------------------------------------------------------------------
def test_06_pdf_content_audit_exists():
    assert _PDF_CONT_AUDIT.exists(), str(_PDF_CONT_AUDIT)
    assert _PDF_CONT_AUDIT.stat().st_size > 512


# ---------------------------------------------------------------------------
# 7. Cross-consistency audit exists
# ---------------------------------------------------------------------------
def test_07_cross_consistency_audit_exists():
    assert _CROSS_AUDIT.exists(), str(_CROSS_AUDIT)
    assert _CROSS_AUDIT.stat().st_size > 512


# ---------------------------------------------------------------------------
# 8. Visual QA index exists
# ---------------------------------------------------------------------------
def test_08_visual_qa_index_exists():
    assert _VI_INDEX.exists(), str(_VI_INDEX)
    assert _VI_INDEX.stat().st_size > 10240


# ---------------------------------------------------------------------------
# 9. Three PDFs exist and are non-empty
# ---------------------------------------------------------------------------
def test_09_three_pdfs_exist_nonempty():
    for p in [_TRAD_PDF, _DET_PDF, _PRO_PDF]:
        assert p.exists(), str(p)
        assert p.stat().st_size > 10240, f"{p.name} too small"


# ---------------------------------------------------------------------------
# 10. Three PDF templates exist
# ---------------------------------------------------------------------------
def test_10_three_pdf_templates_exist():
    for t in [_TRAD_TPL, _DET_TPL, _PRO_TPL]:
        assert t.exists(), str(t)


# ---------------------------------------------------------------------------
# 11. Shared PDF builder exists
# ---------------------------------------------------------------------------
def test_11_pdf_builder_exists():
    assert _BLDR.exists(), str(_BLDR)


# ---------------------------------------------------------------------------
# 12. No internal absolute paths in visual QA index
# ---------------------------------------------------------------------------
def test_12_no_internal_paths_in_index():
    src = _read(_VI_INDEX)
    hits = re.findall(r"[A-Za-z]:\\[^\s\"<>']{5,}", src)
    assert not hits, f"absolute paths in index: {hits[:3]}"


# ---------------------------------------------------------------------------
# 13. No Excel path exposed inside PDF previews/templates
# ---------------------------------------------------------------------------
def test_13_no_excel_path_in_pdf_sources():
    xl_exts = (".xlsx", ".xlsm")
    for pdf, tpl in [(_TRAD_PDF, _TRAD_TPL), (_DET_PDF, _DET_TPL), (_PRO_PDF, _PRO_TPL)]:
        src = _triple(pdf, tpl)
        hits = [e for e in xl_exts if e in src]
        assert not hits, f"{pdf.name}: Excel extension found: {hits}"


# ---------------------------------------------------------------------------
# 14. No automatic Certified output
# ---------------------------------------------------------------------------
_AUTO_CERT = "✓ معتمد"


def test_14_no_auto_certified_output():
    for pdf, tpl in [(_TRAD_PDF, _TRAD_TPL), (_DET_PDF, _DET_TPL), (_PRO_PDF, _PRO_TPL)]:
        assert _AUTO_CERT not in _triple(pdf, tpl), f"{pdf.name}: auto-cert found"


# ---------------------------------------------------------------------------
# 15. No fake signature / license / stamp
# ---------------------------------------------------------------------------
_FAKE = ("signature.png", "license.png", "stamp.png")


def test_15_no_fake_signature_license_stamp():
    for pdf, tpl in [(_TRAD_PDF, _TRAD_TPL), (_DET_PDF, _DET_TPL), (_PRO_PDF, _PRO_TPL)]:
        src = _triple(pdf, tpl)
        hits = [f for f in _FAKE if f in src]
        assert not hits, f"{pdf.name}: fake assets: {hits}"


# ---------------------------------------------------------------------------
# 16. Summary JSON exists
# ---------------------------------------------------------------------------
def test_16_summary_json_exists():
    assert _SUMMARY_JSON.exists(), str(_SUMMARY_JSON)
    assert _SUMMARY_JSON.stat().st_size > 256


# ---------------------------------------------------------------------------
# 17. Overall status is PASS or PARTIAL with blockers documented
# ---------------------------------------------------------------------------
def test_17_overall_status_acceptable():
    data = json.loads(_SUMMARY_JSON.read_text(encoding="utf-8"))
    status = data.get("overall_status", "")
    assert status in ("PASS", "PARTIAL"), f"unexpected status: {status!r}"
    if status == "PARTIAL":
        assert data.get("blockers"), "PARTIAL status must list blockers"
