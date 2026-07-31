# -*- coding: utf-8 -*-
"""
Batch 2 — Certification Pack Merge tests.
25 tests verifying the full merge, regression protection, and governance rules.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import pathlib
import re
import zipfile

import openpyxl
import pytest

# ── Canonical paths ────────────────────────────────────────────────────────────
BASE = pathlib.Path(__file__).resolve().parent.parent.parent
BUILDER_PATH = BASE / "core_engine" / "reports" / "excel_template_driven_builder.py"
BATCH1_PATH = (
    BASE / "core_engine" / "instance" / "manual_review_outputs"
    / "professional_valuation_template_driven_batch1"
    / "excel_outputs" / "template_driven_professional_workbook_batch1.xlsx"
)
CERT_SOURCE_PATH = (
    BASE / "core_engine" / "instance" / "manual_review_outputs"
    / "valuation_certification_readiness_gate"
    / "03_market_certification_readiness_workbook.xlsx"
)
OUTPUT_PATH = (
    BASE / "core_engine" / "instance" / "manual_review_outputs"
    / "professional_valuation_template_driven_batch2"
    / "excel_outputs" / "template_driven_professional_workbook_batch2.xlsx"
)
AUDITS_DIR = (
    BASE / "core_engine" / "instance" / "manual_review_outputs"
    / "professional_valuation_template_driven_batch2" / "audits"
)
PV_OUTPUTS_PATH = BASE / "core_engine" / "professional_valuation_outputs.py"

CERT_REQUIRED_SHEETS = [
    "بيان الامتثال",
    "توقيع واعتماد الخبير",
    "حوكمة مصادر البيانات",
    "قائمة المستندات ومخاطر الاعتماد",
    "حالة الاعتماد والتوصية",
    "نطاق العمل",
    "التوصية النهائية",
    "الإفصاحات المهنية",
    "نطاق الثقة وعدم اليقين",
    "حوكمة المعاملات",
]

_CTX = {
    "request_summary": {
        "client_name": "بنك التجمع الخامس",
        "property_type": "شقة سكنية",
        "property_address": "القاهرة الجديدة",
        "area_sqm": 175,
        "construction_year": 2015,
        "floor_number": 3,
    },
    "method_summary": {
        "cap_rate": 0.09,
        "annual_rent_per_sqm": 420,
    },
    "comparable_summary": {
        "price_per_sqm": 28500,
    },
}


# ── Module + build fixtures ────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def builder_mod():
    spec = importlib.util.spec_from_file_location(
        "excel_template_driven_builder", str(BUILDER_PATH)
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def build_result2(builder_mod):
    result2 = builder_mod.build_certification_merged_xlsx(
        batch1_path=BATCH1_PATH,
        cert_source_path=CERT_SOURCE_PATH,
        ctx=_CTX,
        output_path=OUTPUT_PATH,
    )
    out_dir = OUTPUT_PATH.parent.parent
    b1_inv = builder_mod._template_inventory(BATCH1_PATH)
    cert_inv = builder_mod._template_inventory(CERT_SOURCE_PATH)
    builder_mod.write_batch2_audits(result2, b1_inv, cert_inv, out_dir)
    return result2


@pytest.fixture(scope="module")
def output_wb(build_result2):
    wb = openpyxl.load_workbook(str(OUTPUT_PATH))
    yield wb
    wb.close()


# ── Test 01 — Batch 1 input exists ────────────────────────────────────────────
def test_01_batch1_input_exists():
    assert BATCH1_PATH.is_file(), f"Batch 1 input missing: {BATCH1_PATH}"


# ── Test 02 — Cert source exists ──────────────────────────────────────────────
def test_02_cert_source_exists():
    assert CERT_SOURCE_PATH.is_file(), f"Cert source missing: {CERT_SOURCE_PATH}"


# ── Test 03 — Source hashes unchanged ─────────────────────────────────────────
def test_03_source_hashes_unchanged(build_result2):
    assert build_result2["batch1_sha256_before"] == build_result2["batch1_sha256_after"], \
        "Batch 1 source was mutated during build"
    assert build_result2["cert_sha256_before"] == build_result2["cert_sha256_after"], \
        "Cert source was mutated during build"


# ── Test 04 — Builder exists ──────────────────────────────────────────────────
def test_04_builder_exists():
    assert BUILDER_PATH.is_file(), f"Builder missing: {BUILDER_PATH}"


# ── Test 05 — Batch 2 output exists ──────────────────────────────────────────
def test_05_output_exists(build_result2):
    assert OUTPUT_PATH.is_file(), f"Batch 2 output missing: {OUTPUT_PATH}"


# ── Test 06 — Output opens cleanly ────────────────────────────────────────────
def test_06_output_opens_cleanly():
    wb = openpyxl.load_workbook(str(OUTPUT_PATH))
    assert wb is not None
    wb.close()


# ── Test 07 — Output has exactly 54 sheets ────────────────────────────────────
def test_07_output_has_54_sheets(output_wb):
    assert len(output_wb.sheetnames) == 54, \
        f"Expected 54 sheets, got {len(output_wb.sheetnames)}"


# ── Test 08 — Original 44 sheet names/order unchanged ────────────────────────
def test_08_original_44_sheets_preserved(output_wb):
    wb1 = openpyxl.load_workbook(str(BATCH1_PATH), read_only=True)
    b1_names = list(wb1.sheetnames)
    wb1.close()
    assert list(output_wb.sheetnames)[:44] == b1_names, \
        "First 44 sheet names/order differ from Batch 1"


# ── Test 09 — Ten cert sheets in correct order ────────────────────────────────
def test_09_cert_sheets_in_correct_order(output_wb):
    cert_in_output = list(output_wb.sheetnames)[44:]
    assert cert_in_output == CERT_REQUIRED_SHEETS, (
        f"Cert sheet order incorrect:\n"
        f"  got:      {cert_in_output}\n"
        f"  expected: {CERT_REQUIRED_SHEETS}"
    )


# ── Test 10 — No sheet overwritten or renamed ─────────────────────────────────
def test_10_no_sheet_overwritten(build_result2):
    assert not build_result2["sheet_conflicts"], \
        f"Sheet conflicts: {build_result2['sheet_conflicts']}"
    # Verify all 10 required sheets are reported as copied
    assert build_result2["required_sheets_missing"] == [], \
        f"Missing required sheets: {build_result2['required_sheets_missing']}"


# ── Test 11 — B27:B33 formulas unchanged ──────────────────────────────────────
def test_11_b27_b33_formulas_unchanged(build_result2):
    reg = build_result2.get("batch1_regression", {})
    b_before = reg.get("b27_b33_before", {})
    b_after = reg.get("b27_b33_after", {})
    for coord in ["B27", "B28", "B29", "B30", "B31", "B32", "B33"]:
        before = b_before.get(coord, {}).get("value")
        after = b_after.get(coord, {}).get("value")
        assert before == after, f"{coord} formula changed: {before!r} → {after!r}"


# ── Test 12 — KPI formulas unchanged ──────────────────────────────────────────
def test_12_kpi_formulas_unchanged(build_result2):
    reg = build_result2.get("batch1_regression", {})
    kpi_before = reg.get("kpi_before", {})
    kpi_after = reg.get("kpi_after", {})
    for coord in ["A5", "C5", "G5", "I5"]:
        before = kpi_before.get(coord)
        after = kpi_after.get(coord)
        assert before == after, \
            f"Dashboard KPI {coord} changed: {before!r} → {after!r}"


# ── Test 13 — I1 (dashboard completion formula) unchanged ────────────────────
def test_13_i1_completion_formula_unchanged(build_result2, output_wb):
    reg = build_result2.get("batch1_regression", {})
    kpi_before = reg.get("kpi_before", {})
    kpi_after = reg.get("kpi_after", {})
    i1_before = kpi_before.get("I1")
    i1_after = kpi_after.get("I1")
    assert i1_before == i1_after, \
        f"Dashboard I1 changed: {i1_before!r} → {i1_after!r}"
    # Also verify الافتراضات والمدخلات H1 is still the COUNTA formula
    ws_assump = output_wb["الافتراضات والمدخلات"]
    h1 = ws_assump["H1"].value
    assert h1 and isinstance(h1, str) and h1.startswith("="), \
        f"H1 completion counter is not a formula: {h1!r}"
    assert "COUNTA" in h1.upper(), \
        f"H1 formula does not contain COUNTA: {h1!r}"


# ── Test 14 — Dashboard charts remain ────────────────────────────────────────
def test_14_dashboard_charts_remain(build_result2):
    reg = build_result2.get("batch1_regression", {})
    before = reg.get("chart_count_before", 0)
    after = reg.get("chart_count_after", 0)
    assert after >= before, f"Chart count dropped: {before} → {after}"
    assert after >= 2, f"Expected at least 2 charts, found {after}"


# ── Test 15 — B20 blank when wacc_construction absent ────────────────────────
def test_15_b20_blank_without_wacc_construction(output_wb):
    ws = output_wb["الافتراضات والمدخلات"]
    b20 = ws["B20"].value
    assert b20 is None, (
        f"B20 must be blank (wacc_construction not in ctx), got: {b20!r}"
    )


# ── Test 16 — Missing values not converted to zero ────────────────────────────
def test_16_missing_values_not_zero(output_wb):
    ws = output_wb["الافتراضات والمدخلات"]
    for row in range(4, 27):
        if row == 16:
            continue
        val = ws.cell(row, 2).value
        assert val != 0, (
            f"B{row} is 0 — missing inputs must be blank, never zero"
        )
        assert val != 0.0, (
            f"B{row} is 0.0 — missing inputs must be blank, never zero"
        )


# ── Test 17 — Signature gate remains unsigned ─────────────────────────────────
def test_17_signature_gate_unsigned(output_wb):
    ws = output_wb["توقيع واعتماد الخبير"]
    # B10 (حالة الاعتماد) must not be "معتمد"
    b10 = ws["B10"].value
    assert b10 != "معتمد", f"B10 must not be 'معتمد', got: {b10!r}"
    # B11 (الاعتماد النهائي جاهز؟) must start with "لا" (not ready for certification)
    b11 = str(ws["B11"].value or "")
    assert b11.startswith("لا"), f"B11 must start with 'لا' (not certified), got: {b11!r}"
    # B8 (حالة التوقيع) must contain the governance pending text
    b8 = str(ws["B8"].value or "")
    assert "توقيع" in b8, (
        f"B8 must contain signature pending text, got: {b8!r}"
    )


# ── Test 18 — No fake reviewer / licence / stamp ─────────────────────────────
def test_18_no_fake_reviewer_data(output_wb):
    ws = output_wb["توقيع واعتماد الخبير"]
    FAKE_PAT = re.compile(r"VAL-\d+|RICS-\d+|stamp\.(png|jpg)|TRE-\d+", re.IGNORECASE)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                assert not FAKE_PAT.search(str(cell.value)), \
                    f"Fake pattern at {cell.coordinate}: {cell.value!r}"
    # Appraiser name and licence must remain as-is from cert source
    b3 = ws["B3"].value
    b4 = ws["B4"].value
    assert b3 == "غير متاح ضمن بيانات الطلب", \
        f"Appraiser name must stay 'غير متاح', got: {b3!r}"
    assert b4 == "غير متاح ضمن بيانات الطلب", \
        f"Licence must stay 'غير متاح', got: {b4!r}"


# ── Test 19 — No automatic certified status ───────────────────────────────────
def test_19_no_automatic_certified_status(output_wb):
    ws = output_wb["حالة الاعتماد والتوصية"]
    b3 = ws["B3"].value
    assert b3 == "لا", f"Certification ready must be 'لا', got: {b3!r}"
    b2 = ws["B2"].value
    CERTIFIED_VALS = {"معتمد", "جاهز للاعتماد", "مكتمل", "approved"}
    assert str(b2 or "").strip() not in CERTIFIED_VALS, \
        f"Report status must not indicate certified, got: {b2!r}"


# ── Test 20 — No external workbook links ──────────────────────────────────────
def test_20_no_external_workbook_links(build_result2):
    issues = build_result2.get("workbook_issues", {})
    assert issues.get("external_link_count", 0) == 0, \
        f"External links found: {issues.get('external_link_examples', [])}"


# ── Test 21 — No #REF! formulas ───────────────────────────────────────────────
def test_21_no_ref_errors(build_result2):
    issues = build_result2.get("workbook_issues", {})
    assert issues.get("ref_error_count", 0) == 0, \
        f"#REF! errors: {issues.get('ref_error_examples', [])}"


# ── Test 22 — No VBA stream ───────────────────────────────────────────────────
def test_22_no_vba_stream(build_result2):
    macro = build_result2.get("macro_audit", {})
    assert macro.get("verdict") == "CLEAN", \
        f"VBA found: {macro.get('macro_entries', [])}"


# ── Test 23 — No absolute paths exposed ──────────────────────────────────────
def test_23_no_absolute_paths(build_result2):
    issues = build_result2.get("workbook_issues", {})
    assert issues.get("abs_path_count", 0) == 0, \
        f"Absolute paths in cells: {issues.get('abs_path_examples', [])}"


# ── Test 24 — Copy parity audit exists ────────────────────────────────────────
def test_24_copy_parity_audit_exists():
    audit_path = AUDITS_DIR / "03_certification_sheet_copy_parity.json"
    assert audit_path.is_file(), f"Parity audit missing: {audit_path}"
    data = json.loads(audit_path.read_text(encoding="utf-8"))
    assert data.get("total_sheets_copied") == 10, \
        f"Expected 10 sheets copied, got {data.get('total_sheets_copied')}"
    for entry in data.get("copy_parity", []):
        assert entry["status"] == "PASS", \
            f"Sheet {entry['sheet']} copy status: {entry['status']}"


# ── Test 25 — professional_valuation_outputs.py not modified ─────────────────
def test_25_pv_outputs_not_modified_by_batch2():
    content = PV_OUTPUTS_PATH.read_text(encoding="utf-8")
    assert "build_certification_merged_xlsx" not in content, \
        "professional_valuation_outputs.py was modified — contains batch2 function reference"
    # Verify from_scratch builder is still present (untouched)
    assert "openpyxl.Workbook()" in content, \
        "professional_valuation_outputs.py lost the from-scratch builder"
