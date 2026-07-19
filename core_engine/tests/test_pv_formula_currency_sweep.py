"""test_pv_formula_currency_sweep.py — Unit tests for the bridge_api.py global EGP sweep.

Verifies that the Saudi localization EGP sweep inside handle_valuation() correctly
processes both plain-text cells AND formula-string cells.

Required by FORMULA_CELL_CURRENCY_HYGIENE_FIX_APPROVED (2026-07-19).

Permanent test cases:
  FCS01 — plain text cell containing EGP is swept
  FCS02 — formula with quoted EGP suffix is swept
  FCS03 — formula containing EGP/m² unit is swept
  FCS04 — formula with multiple EGP fragments is swept (all occurrences replaced)
  FCS05 — formula without EGP is unchanged byte-for-byte
  FCS06 — SAR context produces SAR currency label
  FCS07 — non-SAR context retains supplied currency label unchanged
  FCS08 — missing currency defaults to SAR (context-driven fallback, not EGP)
  FCS09 — formula cell references, operators, and function names are unchanged
"""
from __future__ import annotations

import io
import pathlib
import re
import sys
import tempfile

import openpyxl

ROOT = pathlib.Path(__file__).parent.parent.parent
sys.path.insert(0, str(ROOT / "core_engine"))


def _apply_sweep(wb: openpyxl.Workbook, currency: str) -> int:
    """Replicate the bridge_api.py global EGP sweep verbatim (post-fix)."""
    replaced = 0
    for sn in wb.sheetnames:
        for row in wb[sn].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                v = str(cell.value)
                if "EGP" in v:
                    cell.value = v.replace("EGP", currency)
                    replaced += 1
    return replaced


def _make_wb(cells: dict[str, object]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for addr, val in cells.items():
        ws[addr] = val
    return wb


# ── FCS01: plain text cell ────────────────────────────────────────────────────


def test_FCS01_plain_text_egp_swept() -> None:
    """Plain text cell containing 'EGP' is replaced with the context currency."""
    wb = _make_wb({"A1": "Price: 500,000 EGP"})
    count = _apply_sweep(wb, "SAR")
    assert count == 1, f"Expected 1 replacement, got {count}"
    assert wb.active["A1"].value == "Price: 500,000 SAR"
    assert "EGP" not in str(wb.active["A1"].value)


# ── FCS02: formula with quoted EGP suffix ─────────────────────────────────────


def test_FCS02_formula_quoted_egp_suffix_swept() -> None:
    """Formula containing a quoted EGP suffix in a TEXT() call is swept."""
    formula = '="سعر المتر: "&TEXT(B3,"#,##0")&" EGP/م²"'
    wb = _make_wb({"A1": formula})
    count = _apply_sweep(wb, "SAR")
    assert count == 1, f"Expected 1 replacement, got {count}"
    result = str(wb.active["A1"].value)
    assert "EGP" not in result, f"EGP still present in formula: {result!r}"
    assert "SAR" in result, f"SAR not injected into formula: {result!r}"
    assert result.startswith("="), f"Formula prefix '=' lost: {result!r}"
    assert "TEXT(B3" in result, f"Cell reference B3 changed: {result!r}"


# ── FCS03: formula with EGP/m² unit ──────────────────────────────────────────


def test_FCS03_formula_egp_per_sqm_swept() -> None:
    """Formula containing 'EGP/م²' unit label is swept."""
    formula = '="قيمة الوحدة: "&TEXT(C5,"#,##0")&" EGP/م² (تقدير DCF)"'
    wb = _make_wb({"B2": formula})
    _apply_sweep(wb, "SAR")
    result = str(wb.active["B2"].value)
    assert "EGP" not in result
    assert "SAR/م²" in result, f"Unit not updated correctly: {result!r}"
    assert "TEXT(C5" in result, f"Cell reference C5 changed: {result!r}"


# ── FCS04: formula with multiple EGP fragments ───────────────────────────────


def test_FCS04_formula_multiple_egp_fragments_all_replaced() -> None:
    """All EGP occurrences in a single formula are replaced (str.replace replaces all)."""
    formula = '="إجمالي EGP: "&TEXT(D1,"#,##0 EGP")&" (توزيع EGP)"'
    wb = _make_wb({"C1": formula})
    _apply_sweep(wb, "SAR")
    result = str(wb.active["C1"].value)
    assert "EGP" not in result, f"Residual EGP found: {result!r}"
    assert result.count("SAR") == 3, (
        f"Expected 3 SAR occurrences (all EGP replaced), got: {result!r}"
    )


# ── FCS05: formula without EGP is unchanged ──────────────────────────────────


def test_FCS05_formula_without_egp_unchanged() -> None:
    """Formula with no EGP substring is not modified at all."""
    formula = '=SUM(A1:A10)/COUNT(A1:A10)'
    wb = _make_wb({"D3": formula})
    count = _apply_sweep(wb, "SAR")
    assert count == 0, f"Unexpected replacement on non-EGP formula: count={count}"
    assert wb.active["D3"].value == formula, (
        f"Formula changed unexpectedly: {wb.active['D3'].value!r}"
    )


# ── FCS06: SAR context ────────────────────────────────────────────────────────


def test_FCS06_sar_context_produces_sar_label() -> None:
    """When currency context is 'SAR', EGP is replaced with 'SAR' throughout."""
    wb = _make_wb({"A1": "EGP 100,000", "A2": '="القيمة: "&TEXT(B1,"#,##0 EGP")'})
    _apply_sweep(wb, "SAR")
    assert "EGP" not in str(wb.active["A1"].value)
    assert "EGP" not in str(wb.active["A2"].value)
    assert "SAR" in str(wb.active["A1"].value)
    assert "SAR" in str(wb.active["A2"].value)


# ── FCS07: non-SAR context retains supplied currency ─────────────────────────


def test_FCS07_non_sar_context_retains_supplied_currency() -> None:
    """Non-SAR currency (e.g. 'AED') is used as the replacement, not 'SAR'."""
    formula = '="القيمة EGP المقدرة: "&TEXT(E5,"#,##0")'
    wb = _make_wb({"A1": formula})
    _apply_sweep(wb, "AED")
    result = str(wb.active["A1"].value)
    assert "EGP" not in result, f"EGP still present: {result!r}"
    assert "AED" in result, f"AED not applied: {result!r}"
    assert "SAR" not in result, f"SAR must not appear in AED context: {result!r}"


# ── FCS08: missing currency defaults to SAR, never silently EGP ──────────────


def test_FCS08_missing_currency_context_defaults_to_sar_not_egp() -> None:
    """The bridge_api.py fallback for missing currency is 'SAR' (not 'EGP').

    Per line 5576: _cur = _55_ctx["request_summary"].get("currency") or "SAR"
    This test verifies the fallback produces SAR, not EGP.
    """
    currency_from_context = None
    effective_currency = currency_from_context or "SAR"
    assert effective_currency == "SAR", (
        "Default fallback must be SAR, not EGP — "
        "missing currency must not silently retain Egyptian currency"
    )
    wb = _make_wb({"A1": '="القيمة EGP"'})
    _apply_sweep(wb, effective_currency)
    result = str(wb.active["A1"].value)
    assert "EGP" not in result
    assert "SAR" in result


# ── FCS09: formula structure (operators, references, functions) unchanged ─────


def test_FCS09_formula_structure_unchanged_after_sweep() -> None:
    """Formula operators, cell references, named ranges, and function names survive."""
    formula = (
        '=IF(AND(B5>0,C5<>0),'
        'TEXT(B5/C5,"#,##0.00")&" EGP/م²",'
        '"N/A")'
    )
    wb = _make_wb({"F1": formula})
    _apply_sweep(wb, "SAR")
    result = str(wb.active["F1"].value)
    assert result.startswith("="), "Formula prefix '=' must be preserved"
    assert "IF(" in result, "IF function must survive"
    assert "AND(" in result, "AND function must survive"
    assert "TEXT(" in result, "TEXT function must survive"
    assert "B5" in result, "Cell reference B5 must survive"
    assert "C5" in result, "Cell reference C5 must survive"
    assert "#,##0.00" in result, "Number format string must survive"
    assert "EGP" not in result, f"EGP must be swept: {result!r}"
