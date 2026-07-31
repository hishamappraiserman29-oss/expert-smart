"""test_pv_four_workbook_parity.py — Batch 4 final parity tests for professional valuation workbook.

Verifies:
- All 50 original sheets preserved
- All 23 new sheets present (Batch 1: 9, Batch 2: 10, Batch 3: 4)
- Total sheet count = 73
- Assumption bank + cross-reference center in Input Control Panel
- Method Reconciliation formula link in خريطة النتائج
- Gallery has 10 images or documented blockers
- Status-color semantics present in 9 key sheets
- Tab colors set on all new sheets
- Guard text present in new sheets
- No VBA/macro, opens cleanly, no internal paths, no fake official data
"""
import pathlib
import re
import pytest
import openpyxl

ROOT = pathlib.Path(__file__).parent.parent.parent
WB = (
    ROOT
    / "core_engine"
    / "instance"
    / "manual_review_outputs"
    / "professional_valuation_excel_reference_parity"
    / "excel_outputs"
    / "core_valuation_master_workbook_reference_parity.xlsx"
)

ORIGINAL_50 = [
    "Cover", "Sheet Index", "Executive Dashboard", "Input Control Panel",
    "Property Data", "Location Data", "Spatial Analysis", "GIS Location Score",
    "Source Registry", "Data Quality", "Market Evidence", "Comparable Sales",
    "Sales Adjustment Matrix", "Land Sales Comparison", "Land Adjustment Matrix",
    "Land Extraction Method", "Residual Land Method", "Land Reconciliation",
    "Building Cost Breakdown", "Cost Approach", "Depreciation Analysis",
    "Site Improvements", "Income Capitalization", "Rent Comparables",
    "DCF Model", "NPV Analysis", "IRR Analysis", "Payback Analysis",
    "Sale vs Rent", "AVM Summary", "Multiple Regression", "ANN Neural Network",
    "Model Accuracy", "Scenarios", "Sensitivity Matrix", "Risk Matrix",
    "Risk Register", "HBU Analysis", "SWOT", "ESG Sustainability",
    "Standards Readiness", "Certification Readiness", "Method Reconciliation",
    "Final Value", "Charts Dashboard", "Print Report Summary", "Appendix Data",
    "Reviewer Overrides", "Changelog", "Legacy Archive Index",
]

BATCH1_9 = [
    "مقدمة التقرير", "نطاق العمل", "الافتراضات الخاصة والقيود",
    "المستندات والمخاطر", "الفحص القانوني", "بيان الامتثال",
    "خارطة طريق الاعتماد", "توقيع الخبير وبوابة الاعتماد", "حزمة الإيجار",
]

BATCH2_10 = [
    "📊 Hedonic Pricing", "📈 Gordon Growth", "🔁 Repeat Sales",
    "🎲 Monte Carlo", "🏗️ RCNLD", "🏛️ Reproduction Cost",
    "🎯 Expected Utility", "📉 ARIMA", "Real Options", "🔬 طرق إضافية",
]

BATCH3_4 = [
    "🧠 خريطة المنهجية", "🌟 خريطة النتائج",
    "📈 معرض التصوّرات", "🏆 شهادة",
]

NEW_23 = BATCH1_9 + BATCH2_10 + BATCH3_4

STATUS_SHEETS = [
    "Standards Readiness", "Data Quality", "Risk Register", "Risk Matrix",
    "Certification Readiness", "بيان الامتثال", "خارطة طريق الاعتماد",
    "توقيع الخبير وبوابة الاعتماد", "🏆 شهادة",
]


@pytest.fixture(scope="module")
def wb_ro():
    wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
    yield wb
    wb.close()


@pytest.fixture(scope="module")
def wb_formula():
    wb = openpyxl.load_workbook(WB, data_only=False)
    yield wb
    wb.close()


def test_zero_loss_all_50_original_sheets(wb_ro):
    missing = [s for s in ORIGINAL_50 if s not in wb_ro.sheetnames]
    assert missing == [], f"Original sheets missing: {missing}"


def test_four_original_charts_present(wb_ro):
    assert "Charts Dashboard" in wb_ro.sheetnames, "Charts Dashboard sheet missing"
    ws = wb_ro["Charts Dashboard"]
    chart_labels = ["مقارنة المنهجيات", "تفصيل تكلفة المبنى", "تدفقات DCF", "سيناريوهات القيمة"]
    all_vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    found = [lbl for lbl in chart_labels if any(lbl in v for v in all_vals)]
    assert len(found) >= 4, f"Expected 4 original chart entries; found only: {found}"


def test_23_new_sheets_present(wb_ro):
    missing = [s for s in NEW_23 if s not in wb_ro.sheetnames]
    assert missing == [], f"New sheets missing: {missing}"


def test_total_sheet_count_73(wb_ro):
    count = len(wb_ro.sheetnames)
    assert count == 73, f"Expected 73 sheets, got {count}"


def test_assumptions_center_cross_refs(wb_ro):
    assert "Input Control Panel" in wb_ro.sheetnames
    ws = wb_ro["Input Control Panel"]
    all_vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("ASSUMPTION_BANK" in v for v in all_vals), "ASSUMPTION_BANK marker not found"
    assert any("CROSS_REF_CENTER" in v for v in all_vals), "CROSS_REF_CENTER marker not found"


def test_mind_map_results_formula_linked(wb_formula):
    assert "🌟 خريطة النتائج" in wb_formula.sheetnames
    ws = wb_formula["🌟 خريطة النتائج"]
    has_link = any(
        c.value and isinstance(c.value, str) and "'Method Reconciliation'" in c.value
        for row in ws.iter_rows()
        for c in row
        if c.value
    )
    assert has_link, "🌟 خريطة النتائج has no Method Reconciliation formula link"


def test_gallery_has_10_images(wb_formula):
    assert "📈 معرض التصوّرات" in wb_formula.sheetnames
    ws = wb_formula["📈 معرض التصوّرات"]
    summary_val = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and "GALLERY_SUMMARY:" in cell.value:
                summary_val = cell.value
                break
    if summary_val:
        imgs_m = re.search(r"images=(\d+)", summary_val)
        blks_m = re.search(r"blockers=(\d+)", summary_val)
        imgs_n = int(imgs_m.group(1)) if imgs_m else 0
        blks_n = int(blks_m.group(1)) if blks_m else 0
        total = imgs_n + blks_n
        assert total == 10, f"Gallery total not 10: images={imgs_n} blockers={blks_n}"
    else:
        img_count = len(ws._images) if hasattr(ws, "_images") else 0
        assert img_count >= 10, (
            f"Gallery has {img_count} images and no GALLERY_SUMMARY marker"
        )


def test_status_color_semantics_present(wb_ro):
    COLOR_INDICATORS = {"🟢", "🔴", "🟡", "🟠", "✅", "⬜", "⚠", "💜"}
    for sname in STATUS_SHEETS:
        assert sname in wb_ro.sheetnames, f"Status sheet missing: {sname}"
        ws = wb_ro[sname]
        all_vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        has_marker = any("STATUS_COLOR_SEMANTICS" in v for v in all_vals)
        has_indicators = any(
            any(ind in v for ind in COLOR_INDICATORS) for v in all_vals
        )
        assert has_marker or has_indicators, (
            f"Sheet '{sname}' lacks STATUS_COLOR_SEMANTICS marker or color indicators"
        )


def test_new_sheet_tab_colors_present(wb_formula):
    for sname in NEW_23:
        assert sname in wb_formula.sheetnames, f"New sheet missing: {sname}"
        ws = wb_formula[sname]
        tc = ws.sheet_properties.tabColor
        assert tc is not None, f"Sheet '{sname}' has no tab color set"


def test_unavailable_data_guard_present(wb_ro):
    guard = "غير متاح ضمن بيانات الطلب"
    sheets_with_guard = []
    sheets_without = []
    for sname in NEW_23:
        if sname not in wb_ro.sheetnames:
            continue
        ws = wb_ro[sname]
        has = any(
            cell.value and guard in str(cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value
        )
        (sheets_with_guard if has else sheets_without).append(sname)
    pct = len(sheets_with_guard) / len(NEW_23)
    assert pct >= 0.80, (
        f"Only {len(sheets_with_guard)}/{len(NEW_23)} new sheets have guard text. "
        f"Missing: {sheets_without}"
    )


def test_no_macro_xlsx():
    raw = WB.read_bytes()
    assert b"vbaProject" not in raw, "VBA stream detected"
    assert b"xl/vbaProject" not in raw, "VBA project detected"


def test_workbook_opens_clean_openpyxl():
    try:
        wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
        count = len(wb.sheetnames)
        wb.close()
    except Exception as e:
        pytest.fail(f"Workbook failed to open: {e}")
    assert count > 0, "Workbook opened but has no sheets"


def test_no_internal_paths():
    raw = WB.read_bytes()
    raw_str = raw.decode("utf-8", errors="ignore")
    assert r"C:\Users\Lenovo" not in raw_str, "Internal Windows path exposed"
    assert "C:/Users/Lenovo" not in raw_str, "Internal Windows path exposed"


def test_no_fake_official_data():
    raw = WB.read_bytes()
    raw_str = raw.decode("utf-8", errors="ignore")
    fake_flags = [
        "fake_valuer_created=True",
        "fake_stamp_created=True",
        "fake_license_created=True",
        "fake_production_data_created=True",
        "fake_sources_created=True",
    ]
    found = [f for f in fake_flags if f in raw_str]
    assert not found, f"Fake official data flags found: {found}"
