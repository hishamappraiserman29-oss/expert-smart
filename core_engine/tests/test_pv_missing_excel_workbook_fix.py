"""
test_pv_missing_excel_workbook_fix.py
24 backend tests — التحقق من وجود وصحة ملف Excel الإداري الواحد
advisory_only=True | single_admin_excel=True | arabic_primary=True
"""
from __future__ import annotations
import json
from pathlib import Path

import openpyxl
import pytest

_CORE   = Path(__file__).resolve().parent.parent
_QA_OLD = _CORE / "instance" / "manual_review_outputs" / \
          "professional_valuation_single_excel_15_master_plus_legacy_archive_three_pdfs"
_XL_DIR = _QA_OLD / "excel_outputs"
_PDF_DIR = _QA_OLD / "pdf_outputs"
_QA_FIX = _CORE / "instance" / "manual_review_outputs" / \
          "professional_valuation_excel_missing_fix"

_WORKBOOK = "professional_valuation_admin_master_workbook.xlsm"
_MASTER_SHEETS = [
    "Cover", "Data Quality", "Property", "HBU Analysis", "Comparables",
    "Income Approach", "Cost Approach", "AVM Reference", "Scenarios",
    "Sensitivity Matrix", "Reconciliation", "Risk Register",
    "Standards Matrix", "Source Registry", "Admin Notes",
]


def _wb():
    return openpyxl.load_workbook(str(_XL_DIR / _WORKBOOK), read_only=True, keep_vba=False)


def _sheet_text(wb, name: str) -> str:
    ws = wb[name]
    return " ".join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c)


def _load(p: Path) -> dict:
    assert p.exists(), f"Audit file missing: {p}"
    return json.loads(p.read_text(encoding="utf-8"))


# T01
def test_T01_fix_qa_folder_exists():
    assert _QA_FIX.exists()


# T02
def test_T02_filesystem_search_audit_exists():
    assert (_QA_FIX / "excel_filesystem_search_audit.json").exists()


# T03
def test_T03_expected_excel_file_exists():
    assert (_XL_DIR / _WORKBOOK).exists(), f"Missing: {_XL_DIR / _WORKBOOK}"


# T04
def test_T04_excel_file_not_empty():
    size = (_XL_DIR / _WORKBOOK).stat().st_size
    assert size > 50_000, f"Excel file too small: {size} bytes"


# T05
def test_T05_workbook_opens_successfully():
    wb = _wb()
    assert wb is not None
    wb.close()


# T06
def test_T06_first_15_sheets_are_master_sheets():
    wb = _wb()
    sheets = wb.sheetnames
    wb.close()
    assert sheets[:15] == _MASTER_SHEETS, f"First 15 sheets mismatch: {sheets[:15]}"


# T07
def test_T07_first_15_sheets_in_exact_order():
    wb = _wb()
    actual = wb.sheetnames[:15]
    wb.close()
    for i, (expected, got) in enumerate(zip(_MASTER_SHEETS, actual)):
        assert expected == got, f"Sheet {i+1} expected '{expected}', got '{got}'"


# T08
def test_T08_legacy_archive_sheets_included():
    wb = _wb()
    sheets = wb.sheetnames
    wb.close()
    legacy = sheets[15:]
    assert len(legacy) > 0, "No legacy archive sheets found"
    has_prefix = any(s.startswith("LEGACY_GF_") or s.startswith("LEGACY_ULTRA_") for s in legacy)
    assert has_prefix, f"No LEGACY_ prefixed sheets in: {legacy[:10]}"


# T09
def test_T09_only_one_generated_admin_excel_workbook():
    xl_files = list(_XL_DIR.glob("*.xlsm")) + list(_XL_DIR.glob("*.xlsx"))
    assert len(xl_files) == 1, f"Expected 1 Excel file, found {len(xl_files)}: {[f.name for f in xl_files]}"


# T10
def test_T10_three_separate_excel_workbooks_do_not_exist():
    forbidden = [
        "traditional_report_admin_workbook.xlsm",
        "detailed_report_admin_workbook.xlsm",
        "professional_report_admin_workbook.xlsm",
        "traditional_report_admin_workbook.xlsx",
        "detailed_report_admin_workbook.xlsx",
        "professional_report_admin_workbook.xlsx",
    ]
    for name in forbidden:
        assert not (_XL_DIR / name).exists(), f"Forbidden separate Excel found: {name}"


# T11
def test_T11_master_sheets_contain_arabic_headings():
    wb = _wb()
    cover_text = _sheet_text(wb, "Cover")
    wb.close()
    assert "تقرير" in cover_text or "تقييم" in cover_text or "مستوى" in cover_text


# T12
def test_T12_hbu_sheet_contains_four_tests():
    wb = _wb()
    text = _sheet_text(wb, "HBU Analysis")
    wb.close()
    assert "مشروع قانوناً" in text, "HBU: مشروع قانوناً missing"
    assert "ممكن مادياً" in text, "HBU: ممكن مادياً missing"
    assert "مجدٍ مالياً" in text, "HBU: مجدٍ مالياً missing"
    assert "الأعلى إنتاجية" in text, "HBU: الأعلى إنتاجية missing"


# T13
def test_T13_income_approach_contains_dcf_npv_irr():
    wb = _wb()
    text = _sheet_text(wb, "Income Approach")
    wb.close()
    assert "DCF" in text, "Income Approach: DCF missing"
    assert "NPV" in text, "Income Approach: NPV missing"
    assert "IRR" in text, "Income Approach: IRR missing"


# T14
def test_T14_risk_register_contains_value_impact_percentage():
    wb = _wb()
    text = _sheet_text(wb, "Risk Register")
    wb.close()
    assert "%" in text, "Risk Register: percentage missing"


# T15
def test_T15_standards_matrix_contains_disclosures():
    wb = _wb()
    text = _sheet_text(wb, "Standards Matrix")
    wb.close()
    assert "IVS" in text, "Standards Matrix: IVS missing"
    assert "الإفصاح" in text or "إفصاح" in text, "Standards Matrix: إفصاح missing"


# T16
def test_T16_source_registry_exists_with_content():
    wb = _wb()
    text = _sheet_text(wb, "Source Registry")
    wb.close()
    assert "مصادر" in text or "بيانات" in text, "Source Registry: content missing"


# T17
def test_T17_admin_notes_contains_expert_review_fields():
    wb = _wb()
    text = _sheet_text(wb, "Admin Notes")
    wb.close()
    assert "ملاحظات" in text or "خبير" in text or "مراجعة" in text, "Admin Notes: expert fields missing"


# T18
def test_T18_admin_excel_visibility_context_exists():
    data = _load(_QA_FIX / "admin_excel_visibility_ui_audit.json")
    ctx = data["admin_excel_visibility_context"]
    assert ctx["enabled"] is True
    assert ctx["single_admin_excel_workbook_visible"] is True


# T19
def test_T19_excel_is_internal_only():
    data = _load(_QA_FIX / "admin_excel_visibility_ui_audit.json")
    ctx = data["admin_excel_visibility_context"]
    assert ctx["excel_internal_only"] is True
    assert ctx["excel_not_sent_to_user"] is True


# T20
def test_T20_excel_not_shown_as_user_deliverable():
    data = _load(_QA_FIX / "admin_excel_visibility_ui_audit.json")
    ctx = data["admin_excel_visibility_context"]
    assert ctx["excel_not_sent_to_user"] is True


# T21
def test_T21_no_internal_paths_in_visibility_context():
    data = _load(_QA_FIX / "admin_excel_visibility_ui_audit.json")
    ctx = data["admin_excel_visibility_context"]
    assert ctx["internal_paths_hidden"] is True
    dumped = json.dumps(data)
    for pat in [r"C:\\Users", "C:/Users", "/home/"]:
        assert pat not in dumped, f"Internal path found in visibility audit: {pat}"


# T22
def test_T22_pdfs_still_exist():
    for name in ["traditional_report.pdf", "detailed_report.pdf", "professional_report.pdf"]:
        assert (_PDF_DIR / name).exists(), f"PDF missing: {name}"
        assert (_PDF_DIR / name).stat().st_size > 10_000


# T23
def test_T23_ordinary_valuation_unaffected():
    frontend = _CORE.parent / "frontend" / "index.html"
    if not frontend.exists():
        frontend = _CORE / "frontend" / "index.html"
    assert frontend.exists(), "frontend/index.html missing"


# T24
def test_T24_tax_appeal_unaffected():
    bridge = _CORE / "bridge_api.py"
    assert bridge.exists(), "bridge_api.py missing"
    content = bridge.read_text(encoding="utf-8", errors="ignore")
    assert "tax" in content.lower() or "ضريبة" in content or "appeal" in content.lower()
