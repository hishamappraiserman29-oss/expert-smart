"""
test_pv_sim_enhanced_v1.py
Tests for the enhanced simulation report sourcing layer, methods application,
HTML section injection, Excel extension, and end-to-end orchestration.

Tests are structured so every group can pass even when optional dependencies
(Playwright, enrichment providers, live internet) are unavailable.

Run (from repo root):
  python -m pytest core_engine/tests/test_pv_sim_enhanced_v1.py -v
"""
from __future__ import annotations

import pathlib
import sys

import pytest

# ── Path setup ────────────────────────────────────────────────────────────────
_REPO  = pathlib.Path(__file__).resolve().parent.parent.parent
_CORE  = _REPO / "core_engine"
for _p in [str(_CORE), str(_CORE / "reports")]:
    if _p not in sys.path:
        sys.path.insert(0, _p)

# ── Standard QA body (Lusail villa, QAR) ─────────────────────────────────────
QA_BODY = {
    "case_id":           "sim_enhanced_test_001",
    "country_ar":        "قطر",
    "country_code":      "QA",
    "city_ar":           "لوسيل",
    "district_ar":       "فوكس هيلز",
    "asset_type_ar":     "فيلا سكنية",
    "currency_code":     "QAR",
    "simulation_date":   "2026-07-01",
    "land_area_m2":      1200,
    "built_up_area_m2":  850,
    "monthly_rent":      15000,
    "vacancy_rate":      5,
    "collection_loss":   2,
    "op_expenses":       10,
    "maint_reserve":     5,
    "cap_rate":          6.5,
    "land_price":        2500,
    "construction_cost": 2200,
    "phys_depr":         15,
    "func_depr":         5,
    "ext_depr":          3,
    "weight_sales":      50,
    "weight_income":     35,
    "weight_cost":       15,
    "advisory_only":     True,
    "not_real_training": True,
    "fake_signature_created": False,
}


# ════════════════════════════════════════════════════════════════════════════════
# A — Sourcing Layer
# ════════════════════════════════════════════════════════════════════════════════

class TestSourcingLayer:
    """pv_sim_sourced_inputs.source_subject_inputs()"""

    @pytest.fixture(scope="class")
    def sourcing(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        return source_subject_inputs(QA_BODY)

    def test_returns_sourcing_result(self, sourcing):
        from reports.pv_sim_sourced_inputs import SourcingResult
        assert isinstance(sourcing, SourcingResult)

    def test_provenance_table_at_least_6_inputs(self, sourcing):
        assert len(sourcing.provenance_table) >= 6

    def test_every_input_has_input_id(self, sourcing):
        for p in sourcing.provenance_table:
            assert p.input_id, f"Missing input_id: {p}"

    def test_every_input_has_label_ar(self, sourcing):
        for p in sourcing.provenance_table:
            assert p.label_ar, f"Missing label_ar for {p.input_id}"

    def test_every_input_has_source_type(self, sourcing):
        valid = {"mass_appraisal", "web_research", "enrichment_layer", "body_input", "unavailable"}
        for p in sourcing.provenance_table:
            assert p.source_type in valid, f"{p.input_id} has bad source_type: {p.source_type}"

    def test_every_input_has_status(self, sourcing):
        for p in sourcing.provenance_table:
            assert p.status, f"Missing status for {p.input_id}"

    def test_every_input_has_source_tier(self, sourcing):
        for p in sourcing.provenance_table:
            assert p.source_tier, f"Missing source_tier for {p.input_id}"

    def test_unavailable_inputs_have_reconciliation_note(self, sourcing):
        for p in sourcing.provenance_table:
            if p.source_type == "unavailable":
                assert p.reconciliation_note, (
                    f"{p.input_id} is unavailable but has no reason in reconciliation_note"
                )

    def test_subject_inputs_has_required_keys(self, sourcing):
        required = [
            "price_per_m2", "cap_rate", "market_rent_per_m2_monthly",
            "vacancy_rate", "construction_cost_per_m2", "land_price_per_m2",
            "discount_rate", "annual_growth_rate",
            "built_up_area_m2", "land_area_m2",
        ]
        for key in required:
            assert key in sourcing.subject_inputs, f"Missing key in subject_inputs: {key}"

    def test_advisory_only_flag(self, sourcing):
        assert sourcing.advisory_only is True

    def test_external_reference_not_provided(self, sourcing):
        assert sourcing.external_reference_provided is False

    def test_to_dict_is_serialisable(self, sourcing):
        import json
        d = sourcing.to_dict()
        json.dumps(d, ensure_ascii=False)  # must not raise

    def test_no_none_value_without_reason(self, sourcing):
        for p in sourcing.provenance_table:
            if p.value_used is None:
                assert p.reconciliation_note or p.status == "غير متاح", (
                    f"{p.input_id}: value_used is None but no reason given"
                )


# ════════════════════════════════════════════════════════════════════════════════
# B — Draft Governance
# ════════════════════════════════════════════════════════════════════════════════

class TestDraftGovernance:
    """All web/enrichment inputs must be labeled Draft."""

    @pytest.fixture(scope="class")
    def sourcing(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        return source_subject_inputs(QA_BODY)

    def test_web_inputs_are_draft(self, sourcing):
        for p in sourcing.provenance_table:
            if p.source_type == "web_research":
                assert "Draft" in p.status, (
                    f"{p.input_id} is web_research but status is not Draft: {p.status}"
                )

    def test_enrichment_inputs_are_draft(self, sourcing):
        for p in sourcing.provenance_table:
            if p.source_type == "enrichment_layer":
                assert "Draft" in p.status, (
                    f"{p.input_id} is enrichment_layer but status not Draft: {p.status}"
                )

    def test_no_api_keys_in_source_uris(self, sourcing):
        for p in sourcing.provenance_table:
            uri = p.source_uri.lower()
            for keyword in ("api_key=", "apikey=", "token=", "key=", "secret="):
                assert keyword not in uri, (
                    f"{p.input_id}: API key exposed in source_uri"
                )

    def test_no_api_keys_in_source_log(self, sourcing):
        for entry in sourcing.source_log:
            uri = str(entry.get("url") or entry.get("source_uri") or "").lower()
            for keyword in ("api_key=", "apikey=", "token=", "key=", "secret="):
                assert keyword not in uri, "API key exposed in source_log URI"

    def test_no_local_paths_in_source_log(self, sourcing):
        for entry in sourcing.source_log:
            uri = str(entry.get("url") or entry.get("source_uri") or "")
            assert not uri.startswith("file://"), f"Local path in source_log: {uri}"
            assert ":\\" not in uri, f"Windows path in source_log: {uri}"

    def test_external_reference_flag_false(self, sourcing):
        assert sourcing.external_reference_provided is False

    def test_reference_source_not_external(self, sourcing):
        assert sourcing.reference_source != "external"

    def test_mass_appraisal_inputs_are_certified(self, sourcing):
        for p in sourcing.provenance_table:
            if p.source_type == "mass_appraisal":
                assert p.status == "Certified", (
                    f"{p.input_id} is mass_appraisal but status is not Certified: {p.status}"
                )


# ════════════════════════════════════════════════════════════════════════════════
# C — Sourced Methods
# ════════════════════════════════════════════════════════════════════════════════

class TestSourcedMethods:
    """apply_methods_to_subject — 4 methods computed."""

    @pytest.fixture(scope="class")
    def methods_result(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
        )
        sourcing = source_subject_inputs(QA_BODY)
        inputs   = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights  = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        shadow_results = [sh_s, sh_i, sh_c, sh_d, sh_r]
        return apply_methods_to_subject(sourcing, QA_BODY, scenarios, shadow_results, "QAR")

    def test_returns_methods_result(self, methods_result):
        from reports.pv_sim_market_methods import MethodsResult
        assert isinstance(methods_result, MethodsResult)

    def test_four_methods_in_result(self, methods_result):
        assert len(methods_result.methods) == 4

    def test_method_ids_correct(self, methods_result):
        ids = {m.method_id for m in methods_result.methods}
        expected = {"sales_comparison", "income_capitalisation", "cost_approach", "dcf"}
        assert ids == expected

    def test_computed_values_positive_when_available(self, methods_result):
        for m in methods_result.methods:
            if m.computed_value is not None:
                assert m.computed_value > 0, f"{m.method_id}: computed_value <= 0"

    def test_calculation_steps_non_empty(self, methods_result):
        for m in methods_result.methods:
            if m.computed_value is not None:
                assert len(m.calculation_steps) > 0, (
                    f"{m.method_id} has computed value but no calculation_steps"
                )

    def test_input_prov_ids_non_empty(self, methods_result):
        for m in methods_result.methods:
            assert len(m.input_prov_ids) > 0, (
                f"{m.method_id} has no input_prov_ids"
            )

    def test_governance_flags_on_methods(self, methods_result):
        for m in methods_result.methods:
            assert m.advisory_only is True
            assert m.certification_ready is False
            assert m.fake_signature_created is False

    def test_governance_flags_on_result(self, methods_result):
        assert methods_result.advisory_only is True
        assert methods_result.certification_ready is False
        assert methods_result.fake_signature_created is False

    def test_reconciliation_present(self, methods_result):
        rec = methods_result.reconciliation
        assert isinstance(rec, dict)
        assert "reconciled_value" in rec

    def test_comparison_table_length_matches_methods(self, methods_result):
        assert len(methods_result.comparison_table) == len(methods_result.methods)

    def test_to_dict_serialisable(self, methods_result):
        import json
        json.dumps(methods_result.to_dict(), ensure_ascii=False)

    def test_rag_status_valid(self, methods_result):
        valid_rag = {"🟢", "🟡", "🔴", "⚪"}
        for m in methods_result.methods:
            assert m.rag_status in valid_rag, (
                f"{m.method_id}: unexpected rag_status: {m.rag_status}"
            )


# ════════════════════════════════════════════════════════════════════════════════
# D — Comparison Table
# ════════════════════════════════════════════════════════════════════════════════

class TestComparisonTable:
    """Comparison table structure and diff calculation."""

    @pytest.fixture(scope="class")
    def comparison_table(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        result = apply_methods_to_subject(
            sourcing, QA_BODY, scenarios, [sh_s, sh_i, sh_c, sh_d, sh_r], "QAR"
        )
        return result.comparison_table

    def test_has_four_rows(self, comparison_table):
        assert len(comparison_table) == 4

    def test_required_keys_per_row(self, comparison_table):
        required = [
            "method_id", "method_name_ar", "simulation_value",
            "computed_value", "diff_amount", "diff_pct", "rag_status",
            "input_source_summary", "input_prov_ids",
        ]
        for row in comparison_table:
            for k in required:
                assert k in row, f"Missing key '{k}' in comparison_table row {row.get('method_id')}"

    def test_diff_pct_consistent_with_values(self, comparison_table):
        for row in comparison_table:
            cv = row.get("computed_value")
            sv = row.get("simulation_value")
            dp = row.get("diff_pct")
            if cv and sv and dp is not None:
                expected = round((cv - sv) / sv * 100, 2)
                assert abs(dp - expected) < 0.5, (
                    f"{row['method_id']}: diff_pct={dp} but expected {expected}"
                )

    def test_rag_status_present_all_rows(self, comparison_table):
        for row in comparison_table:
            assert row.get("rag_status"), f"Missing rag_status in {row.get('method_id')}"


# ════════════════════════════════════════════════════════════════════════════════
# E — HTML Sections (user)
# ════════════════════════════════════════════════════════════════════════════════

class TestHTMLSectionsUser:
    """User HTML must have provenance + methods + comparison; NO source log."""

    @pytest.fixture(scope="class")
    def user_html(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from reports.pv_sim_enhanced_html import inject_enhanced_sections
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
            _build_comparison_table, _build_simulation_html_v2,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        shadow_results   = [sh_s, sh_i, sh_c, sh_d, sh_r]
        comparison_table = _build_comparison_table(scenarios, shadow_results)
        methods_result   = apply_methods_to_subject(
            sourcing, QA_BODY, scenarios, shadow_results, "QAR"
        )
        meta = {
            "case_id": "test001", "simulation_date": "2026-07-01",
            "asset_type_ar": "فيلا", "country_ar": "قطر",
            "city_ar": "لوسيل", "district_ar": "فوكس هيلز",
            "land_area_m2": 1200, "built_up_area_m2": 850,
            "weight_sales": 50, "weight_income": 35, "weight_cost": 15,
        }
        base_html = _build_simulation_html_v2(
            meta, scenarios, shadow_results, comparison_table, "user", "QAR"
        )
        return inject_enhanced_sections(base_html, sourcing, methods_result, "user", "QAR")

    def test_html_non_empty(self, user_html):
        assert len(user_html) > 2000

    def test_provenance_section_present(self, user_html):
        assert 'id="enhanced_11"' in user_html

    def test_applied_methods_section_present(self, user_html):
        assert 'id="enhanced_12"' in user_html

    def test_comparison_table_section_present(self, user_html):
        assert 'id="enhanced_13"' in user_html

    def test_source_log_absent_in_user_html(self, user_html):
        assert 'id="enhanced_14"' not in user_html, "Source log must not appear in user HTML"

    def test_governance_badge_present(self, user_html):
        assert "advisory_only=True" in user_html

    def test_advisory_watermark_present(self, user_html):
        assert "استرشادي" in user_html

    def test_no_local_path_in_user_html(self, user_html):
        assert "file://" not in user_html, "Local file:// path exposed in user HTML"
        assert "C:\\" not in user_html, "Windows path exposed in user HTML"

    def test_body_close_tag_present(self, user_html):
        assert "</body>" in user_html


# ════════════════════════════════════════════════════════════════════════════════
# F — HTML Sections (admin)
# ════════════════════════════════════════════════════════════════════════════════

class TestHTMLSectionsAdmin:
    """Admin HTML must have all 4 enhanced sections including source log."""

    @pytest.fixture(scope="class")
    def admin_html(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from reports.pv_sim_enhanced_html import inject_enhanced_sections
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
            _build_comparison_table, _build_simulation_html_v2,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        shadow_results   = [sh_s, sh_i, sh_c, sh_d, sh_r]
        comparison_table = _build_comparison_table(scenarios, shadow_results)
        methods_result   = apply_methods_to_subject(
            sourcing, QA_BODY, scenarios, shadow_results, "QAR"
        )
        meta = {
            "case_id": "test001", "simulation_date": "2026-07-01",
            "asset_type_ar": "فيلا", "country_ar": "قطر",
            "city_ar": "لوسيل", "district_ar": "فوكس هيلز",
            "land_area_m2": 1200, "built_up_area_m2": 850,
            "weight_sales": 50, "weight_income": 35, "weight_cost": 15,
        }
        base_html = _build_simulation_html_v2(
            meta, scenarios, shadow_results, comparison_table, "admin", "QAR"
        )
        return inject_enhanced_sections(base_html, sourcing, methods_result, "admin", "QAR")

    def test_admin_html_non_empty(self, admin_html):
        assert len(admin_html) > 3000

    def test_provenance_section_present(self, admin_html):
        assert 'id="enhanced_11"' in admin_html

    def test_applied_methods_section_present(self, admin_html):
        assert 'id="enhanced_12"' in admin_html

    def test_comparison_table_section_present(self, admin_html):
        assert 'id="enhanced_13"' in admin_html

    def test_source_log_present_in_admin_html(self, admin_html):
        assert 'id="enhanced_14"' in admin_html, "Source log must be present in admin HTML"

    def test_admin_watermark_present(self, admin_html):
        assert "للمراجع المعتمد فقط" in admin_html

    def test_no_local_path_in_admin_html(self, admin_html):
        assert "file://" not in admin_html
        assert "C:\\" not in admin_html

    def test_governance_flags_in_admin_html(self, admin_html):
        assert "advisory_only=True" in admin_html
        assert "certification_ready=False" in admin_html


# ════════════════════════════════════════════════════════════════════════════════
# G — Excel Enhanced Sheets
# ════════════════════════════════════════════════════════════════════════════════

class TestExcelEnhancedSheets:
    """add_enhanced_sheets() adds 2 sheets + charts to existing workbook."""

    @pytest.fixture(scope="class")
    def workbook_and_result(self, tmp_path_factory):
        tmp = tmp_path_factory.mktemp("excel_test")
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
            _build_comparison_table, _build_simulation_excel_v2,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        shadow_results   = [sh_s, sh_i, sh_c, sh_d, sh_r]
        comparison_table = _build_comparison_table(scenarios, shadow_results)
        methods_result   = apply_methods_to_subject(
            sourcing, QA_BODY, scenarios, shadow_results, "QAR"
        )
        meta = {
            "case_id": "test001", "simulation_date": "2026-07-01",
            "asset_type_ar": "فيلا", "country_ar": "قطر",
            "city_ar": "لوسيل", "district_ar": "فوكس هيلز",
            "land_area_m2": 1200, "built_up_area_m2": 850,
            "weight_sales": 50, "weight_income": 35, "weight_cost": 15,
        }
        fname = "test_enhanced.xlsx"
        ok = _build_simulation_excel_v2(
            meta, scenarios, shadow_results, comparison_table, tmp, fname, "QAR"
        )
        xl_path = tmp / fname
        if not (ok and xl_path.exists()):
            pytest.skip("Base Excel generation not available")

        from openpyxl import load_workbook
        from reports.pv_sim_enhanced_excel import add_enhanced_sheets
        wb = load_workbook(str(xl_path))
        wb = add_enhanced_sheets(wb, sourcing.provenance_table, methods_result)
        wb.save(str(xl_path))

        wb2 = load_workbook(str(xl_path))
        return wb2, methods_result

    def test_provenance_sheet_exists(self, workbook_and_result):
        wb, _ = workbook_and_result
        assert "المصدرية" in wb.sheetnames

    def test_comparison_sheet_exists(self, workbook_and_result):
        wb, _ = workbook_and_result
        assert "مقارنة_مصدرية" in wb.sheetnames

    def test_total_sheets_at_least_12(self, workbook_and_result):
        wb, _ = workbook_and_result
        assert len(wb.sheetnames) >= 12

    def test_provenance_sheet_has_data_rows(self, workbook_and_result):
        wb, _ = workbook_and_result
        ws = wb["المصدرية"]
        assert ws.max_row >= 3, "Provenance sheet has too few rows"

    def test_comparison_sheet_has_data_rows(self, workbook_and_result):
        wb, _ = workbook_and_result
        ws = wb["مقارنة_مصدرية"]
        assert ws.max_row >= 3, "Comparison sheet has too few rows"

    def test_workbook_has_charts(self, workbook_and_result):
        wb, _ = workbook_and_result
        total_charts = sum(len(ws._charts) for ws in wb.worksheets)
        assert total_charts >= 2, f"Expected >= 2 charts, found {total_charts}"

    def test_no_local_paths_in_provenance_sheet(self, workbook_and_result):
        wb, _ = workbook_and_result
        ws = wb["المصدرية"]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    assert "file://" not in cell, f"file:// path in provenance sheet: {cell}"
                    assert ":\\" not in cell or "محجوب" in cell, (
                        f"Windows path exposed in provenance sheet: {cell}"
                    )


# ════════════════════════════════════════════════════════════════════════════════
# H — Governance Flags
# ════════════════════════════════════════════════════════════════════════════════

class TestGovernanceFlags:
    """All governance flags must be correct throughout."""

    def test_sourcing_result_governance(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        sr = source_subject_inputs(QA_BODY)
        assert sr.advisory_only is True
        assert sr.external_reference_provided is False

    def test_methods_result_governance(self):
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        mr = apply_methods_to_subject(
            sourcing, QA_BODY, scenarios, [sh_s, sh_i, sh_c, sh_d, sh_r], "QAR"
        )
        assert mr.advisory_only is True
        assert mr.certification_ready is False
        assert mr.fake_signature_created is False

    def test_orchestrator_governance_flags(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report, _SAFETY
        assert _SAFETY["advisory_only"] is True
        assert _SAFETY["certification_ready"] is False
        assert _SAFETY["fake_signature_created"] is False
        assert _SAFETY["non_certified"] is True
        assert _SAFETY["not_real_training"] is True


# ════════════════════════════════════════════════════════════════════════════════
# I — Audience Separation
# ════════════════════════════════════════════════════════════════════════════════

class TestAudienceSeparation:
    """Source log absent in user output; Excel never served to user."""

    def _make_html(self, audience: str) -> str:
        from reports.pv_sim_sourced_inputs import source_subject_inputs
        from reports.pv_sim_market_methods import apply_methods_to_subject
        from reports.pv_sim_enhanced_html import inject_enhanced_sections
        from pv_report_simulation_endpoint import (
            _extract_inputs, _run_three_scenarios,
            _shadow_sales_comparison, _shadow_income_approach,
            _shadow_cost_approach, _shadow_dcf, _shadow_reconciliation,
            _build_comparison_table, _build_simulation_html_v2,
        )
        sourcing  = source_subject_inputs(QA_BODY)
        inputs    = _extract_inputs(QA_BODY)
        scenarios = _run_three_scenarios(inputs)
        weights   = {"sales": 50, "income": 35, "cost": 15}
        sh_s  = _shadow_sales_comparison(inputs, "QAR")
        sh_i  = _shadow_income_approach(inputs, "QAR")
        sh_c  = _shadow_cost_approach(inputs, "QAR")
        sh_d  = _shadow_dcf(inputs, "QAR")
        sh_r  = _shadow_reconciliation(sh_s, sh_i, sh_c, weights, "QAR")
        shadow_results   = [sh_s, sh_i, sh_c, sh_d, sh_r]
        ct = _build_comparison_table(scenarios, shadow_results)
        mr = apply_methods_to_subject(sourcing, QA_BODY, scenarios, shadow_results, "QAR")
        meta = {
            "case_id": "aud_test", "simulation_date": "2026-07-01",
            "asset_type_ar": "فيلا", "country_ar": "قطر",
            "city_ar": "لوسيل", "district_ar": "فوكس هيلز",
            "land_area_m2": 1200, "built_up_area_m2": 850,
            "weight_sales": 50, "weight_income": 35, "weight_cost": 15,
        }
        base = _build_simulation_html_v2(meta, scenarios, shadow_results, ct, audience, "QAR")
        return inject_enhanced_sections(base, sourcing, mr, audience, "QAR")

    def test_source_log_absent_in_user(self):
        html = self._make_html("user")
        assert 'id="enhanced_14"' not in html

    def test_source_log_present_in_admin(self):
        html = self._make_html("admin")
        assert 'id="enhanced_14"' in html

    def test_admin_badge_absent_in_user_html(self):
        html = self._make_html("user")
        assert "للمراجع الداخلي فقط" not in html

    def test_admin_badge_present_in_admin_html(self):
        html = self._make_html("admin")
        assert "للمراجع الداخلي فقط" in html


# ════════════════════════════════════════════════════════════════════════════════
# J — InputProvenance model
# ════════════════════════════════════════════════════════════════════════════════

class TestInputProvenanceModel:
    """Unit tests for InputProvenance dataclass."""

    def test_to_dict_has_all_fields(self):
        from reports.pv_sim_sourced_inputs import InputProvenance
        p = InputProvenance(
            input_id="x", label_ar="تسمية", value_used=100.0, unit="ريال",
            source_type="mass_appraisal", source_name="AVM",
            source_uri="internal:x", retrieved_at="2026-01-01T00:00:00Z",
            confidence_score=80.0, source_tier="Tier-1 محوكم",
            status="Certified",
        )
        d = p.to_dict()
        for k in [
            "input_id", "label_ar", "value_used", "unit",
            "source_type", "source_name", "source_uri",
            "retrieved_at", "confidence_score", "source_tier",
            "status", "reconciliation_note",
        ]:
            assert k in d, f"Missing key {k} in InputProvenance.to_dict()"

    def test_default_reconciliation_note_empty(self):
        from reports.pv_sim_sourced_inputs import InputProvenance
        p = InputProvenance(
            input_id="y", label_ar="اختبار", value_used=None, unit="",
            source_type="unavailable", source_name="—", source_uri="",
            retrieved_at="", confidence_score=0.0,
            source_tier="N/A", status="غير متاح",
        )
        assert p.reconciliation_note == ""
        assert p.mass_appraisal_value is None
        assert p.research_value is None


# ════════════════════════════════════════════════════════════════════════════════
# K — End-to-end smoke test (skipped if Playwright unavailable)
# ════════════════════════════════════════════════════════════════════════════════

class TestEndToEndSmoke:
    """Full orchestrator — skipped gracefully when Playwright/deps missing."""

    def test_orchestrator_returns_dict(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        result = generate_enhanced_simulation_report(
            body=QA_BODY,
            outputs_dir=tmp_path / "out",
            run_id="TEST001",
            currency="QAR",
        )
        assert isinstance(result, dict)

    def test_orchestrator_governance_in_result(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        result = generate_enhanced_simulation_report(
            body=QA_BODY,
            outputs_dir=tmp_path / "out2",
            run_id="TEST002",
            currency="QAR",
        )
        gov = result.get("governance", {})
        assert gov.get("advisory_only") is True
        assert gov.get("certification_ready") is False
        assert gov.get("fake_signature_created") is False

    def test_orchestrator_provenance_table_in_result(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        result = generate_enhanced_simulation_report(
            body=QA_BODY,
            outputs_dir=tmp_path / "out3",
            run_id="TEST003",
            currency="QAR",
        )
        pt = result.get("provenance_table", [])
        assert isinstance(pt, list)
        assert len(pt) >= 6

    def test_orchestrator_html_files_exist_if_no_errors(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        result = generate_enhanced_simulation_report(
            body=QA_BODY,
            outputs_dir=tmp_path / "out4",
            run_id="TEST004",
            currency="QAR",
        )
        if result.get("errors"):
            pytest.skip(f"Orchestrator errors — skipping file checks: {result['errors']}")
        artifacts = result.get("artifacts", {})
        for key in ("user_html", "admin_html"):
            p = artifacts.get(key)
            assert p and pathlib.Path(p).exists(), f"{key} file missing"

    def test_orchestrator_no_local_paths_in_html(self, tmp_path):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        result = generate_enhanced_simulation_report(
            body=QA_BODY,
            outputs_dir=tmp_path / "out5",
            run_id="TEST005",
            currency="QAR",
        )
        artifacts = result.get("artifacts", {})
        for key in ("user_html", "admin_html"):
            p = artifacts.get(key)
            if p and pathlib.Path(p).exists():
                content = pathlib.Path(p).read_text(encoding="utf-8")
                assert "file://" not in content, f"file:// path in {key}"
                assert "C:\\" not in content, f"Windows path in {key}"


# ════════════════════════════════════════════════════════════════════════════════
# L — Enhanced Visual QA  (all 5 formats, single data source)
# ════════════════════════════════════════════════════════════════════════════════

class TestEnhancedVisualQA:
    """
    Visual QA for all 5 enhanced simulation report formats (QA002 Lusail villa, QAR).
    Single data source: generate_enhanced_simulation_report().
    Playwright / PyMuPDF tests skip gracefully when not installed.
    """

    @pytest.fixture(scope="class")
    def enhanced_artifacts(self, tmp_path_factory):
        from reports.pv_sim_enhanced_report import generate_enhanced_simulation_report
        out = tmp_path_factory.mktemp("evqa_out")
        return generate_enhanced_simulation_report(
            body=QA_BODY, outputs_dir=out, run_id="EVQA001", currency="QAR"
        )

    @pytest.fixture(scope="class")
    def vqa_report(self, enhanced_artifacts, tmp_path_factory):
        import sys as _sys
        arts = enhanced_artifacts.get("artifacts", {})
        required = ["user_html", "user_pdf", "admin_html", "admin_pdf", "admin_xlsx"]
        missing = [k for k in required if not arts.get(k) or not pathlib.Path(arts[k]).exists()]
        if missing:
            pytest.skip(f"Artifacts missing — VQA skipped: {missing}")
        _sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
        from pv_simulation_visual_qa import run_simulation_visual_qa
        qa_out = tmp_path_factory.mktemp("evqa_screenshots")
        return run_simulation_visual_qa(
            user_html=arts["user_html"],
            user_pdf=arts["user_pdf"],
            admin_html=arts["admin_html"],
            admin_pdf=arts["admin_pdf"],
            admin_xlsx=arts["admin_xlsx"],
            out_dir=qa_out,
            case_id="simulation_QA002_enhanced",
        )

    # ── HTML structure ─────────────────────────────────────────────────────────

    def test_user_html_sections_11_12_13_present(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        for sec in ["11", "12", "13"]:
            assert f'id="enhanced_{sec}"' in text, f"Section {sec} missing from user HTML"

    def test_user_html_no_section_14(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert 'id="enhanced_14"' not in text, "Source log section 14 must NOT be in user HTML"

    def test_admin_html_sections_11_to_14_present(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        for sec in ["11", "12", "13", "14"]:
            assert f'id="enhanced_{sec}"' in text, f"Section {sec} missing from admin HTML"

    def test_user_html_governance_badge(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert "advisory_only=True" in text
        assert "استرشادي" in text
        assert "certification_ready=False" in text

    def test_admin_html_governance_badge(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert "advisory_only=True" in text
        assert "للمراجع المعتمد فقط" in text

    def test_no_local_paths_user_html(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        for marker in ["file:///", "C:\\", "C:/"]:
            assert marker not in text, f"Local path {marker!r} exposed in user HTML"

    def test_no_local_paths_admin_html(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        for marker in ["file:///", "C:\\", "C:/"]:
            assert marker not in text, f"Local path {marker!r} exposed in admin HTML"

    def test_draft_label_in_user_html(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert "Draft" in text or "مسح مبدئي" in text, "Draft label missing from user HTML"

    def test_value_parity_user_admin_html(self, enhanced_artifacts):
        """Reconciled value must be identical in both user and admin HTML."""
        arts = enhanced_artifacts.get("artifacts", {})
        up, ap = arts.get("user_html"), arts.get("admin_html")
        if not (up and pathlib.Path(up).exists() and ap and pathlib.Path(ap).exists()):
            pytest.skip("HTML files not generated")
        mr = enhanced_artifacts.get("methods_result") or {}
        if isinstance(mr, dict):
            rec_val = mr.get("reconciliation", {}).get("reconciled_value")
        else:
            rec_val = getattr(getattr(mr, "reconciliation", {}), "get", lambda k: None)("reconciled_value")
        if rec_val is None:
            pytest.skip("No reconciled value available")
        rec_str = f"{rec_val:,.0f}"
        user_text  = pathlib.Path(up).read_text(encoding="utf-8")
        admin_text = pathlib.Path(ap).read_text(encoding="utf-8")
        assert rec_str in user_text,  f"Reconciled {rec_str} missing from user HTML"
        assert rec_str in admin_text, f"Reconciled {rec_str} missing from admin HTML"

    # ── Audience separation ────────────────────────────────────────────────────

    def test_no_admin_watermark_in_user_html(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert "للمراجع المعتمد فقط" not in text, "Admin watermark leaked into user HTML"

    def test_admin_watermark_in_admin_html(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_html")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_html not generated")
        text = pathlib.Path(p).read_text(encoding="utf-8")
        assert "للمراجع المعتمد فقط" in text, "Admin watermark missing from admin HTML"

    # ── PDF ─────────────────────────────────────────────────────────────────────

    def test_user_pdf_valid_header(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_pdf")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_pdf not generated")
        assert pathlib.Path(p).read_bytes()[:4] == b"%PDF", "User PDF has invalid header"

    def test_admin_pdf_valid_header(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_pdf")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_pdf not generated")
        assert pathlib.Path(p).read_bytes()[:4] == b"%PDF", "Admin PDF has invalid header"

    def test_user_pdf_size_playwright(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("user_pdf")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("user_pdf not generated")
        size = pathlib.Path(p).stat().st_size
        assert size > 100_000, f"User PDF too small ({size} B) — likely stub not Playwright"

    def test_admin_pdf_size_playwright(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_pdf")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_pdf not generated")
        size = pathlib.Path(p).stat().st_size
        assert size > 100_000, f"Admin PDF too small ({size} B) — likely stub not Playwright"

    def test_user_pdf_all_pages_have_text(self, vqa_report):
        r = vqa_report.get("user_pdf", {})
        if r.get("pass") is None:
            pytest.skip("PyMuPDF not available")
        blank = [pg for pg in r.get("pages", []) if not pg.get("has_text")]
        assert blank == [], f"User PDF blank pages: {blank}"

    def test_admin_pdf_all_pages_have_text(self, vqa_report):
        r = vqa_report.get("admin_pdf", {})
        if r.get("pass") is None:
            pytest.skip("PyMuPDF not available")
        blank = [pg for pg in r.get("pages", []) if not pg.get("has_text")]
        assert blank == [], f"Admin PDF blank pages: {blank}"

    def test_pdf_no_local_paths(self, vqa_report):
        for key in ("user_pdf", "admin_pdf"):
            r = vqa_report.get(key, {})
            if r.get("pass") is None:
                continue
            exposed = [pg for pg in r.get("pages", []) if pg.get("local_path_exposed")]
            assert exposed == [], f"{key} pages expose local paths: {exposed}"

    # ── Playwright HTML VQA ────────────────────────────────────────────────────

    def test_vqa_user_html_no_blank_sections(self, vqa_report):
        r = vqa_report.get("user_html", {})
        if r.get("pass") is None:
            pytest.skip("Playwright not available")
        blank = [s for s in r.get("sections", []) if s.get("blank")]
        assert blank == [], f"Blank user HTML sections: {blank}"

    def test_vqa_admin_html_no_blank_sections(self, vqa_report):
        r = vqa_report.get("admin_html", {})
        if r.get("pass") is None:
            pytest.skip("Playwright not available")
        blank = [s for s in r.get("sections", []) if s.get("blank")]
        assert blank == [], f"Blank admin HTML sections: {blank}"

    def test_vqa_user_html_section_count_ge_13(self, vqa_report):
        r = vqa_report.get("user_html", {})
        if r.get("pass") is None:
            pytest.skip("Playwright not available")
        count = len(r.get("sections", []))
        assert count >= 13, f"Expected ≥13 sections in user HTML, got {count}"

    def test_vqa_admin_html_section_count_ge_14(self, vqa_report):
        r = vqa_report.get("admin_html", {})
        if r.get("pass") is None:
            pytest.skip("Playwright not available")
        count = len(r.get("sections", []))
        assert count >= 14, f"Expected ≥14 sections in admin HTML, got {count}"

    # ── Excel ──────────────────────────────────────────────────────────────────

    def test_admin_xlsx_valid_header(self, enhanced_artifacts):
        p = enhanced_artifacts.get("artifacts", {}).get("admin_xlsx")
        if not p or not pathlib.Path(p).exists():
            pytest.skip("admin_xlsx not generated")
        assert pathlib.Path(p).read_bytes()[:4] == b"PK\x03\x04", "Invalid XLSX header"

    def test_admin_xlsx_sheet_count_ge_12(self, vqa_report):
        r = vqa_report.get("admin_excel_structural", {})
        if r.get("pass") is None:
            pytest.skip("openpyxl not available")
        count = r.get("sheet_count", 0)
        assert count >= 12, f"Expected ≥12 sheets, got {count}"

    def test_admin_xlsx_enhanced_sheets_present(self, vqa_report):
        r = vqa_report.get("admin_excel_structural", {})
        if r.get("pass") is None:
            pytest.skip("openpyxl not available")
        sheets = r.get("sheet_names", [])
        assert "المصدرية" in sheets,      "Provenance sheet 'المصدرية' missing from Excel"
        assert "مقارنة_مصدرية" in sheets, "Comparison sheet 'مقارنة_مصدرية' missing from Excel"

    def test_admin_xlsx_charts_ge_3(self, vqa_report):
        r = vqa_report.get("admin_excel_structural", {})
        if r.get("pass") is None:
            pytest.skip("openpyxl not available")
        count = r.get("chart_count", 0)
        assert count >= 3, f"Expected ≥3 charts (base + 2 enhanced), got {count}"

    def test_admin_xlsx_no_local_paths(self, vqa_report):
        r = vqa_report.get("admin_excel_structural", {})
        if r.get("pass") is None:
            pytest.skip("openpyxl not available")
        issues = [i for i in r.get("issues", []) if "local path" in i.lower()]
        assert issues == [], f"Excel local path issues: {issues}"
