"""
Regression baseline tests for the single-property Excel report builder.

Wave R1 — Golden Baselines (baseline/safety-test only; NOT the Requirements Matrix).

These tests build real .xlsx files with a frozen report_date, extract all
non-empty cell values, and compare against golden snapshots stored in
core_engine/tests/fixtures/.

To regenerate snapshots after an intentional change:
    UPDATE_SNAPSHOTS=1 python -m pytest core_engine/tests/test_report_baseline.py -q

Normal run (asserts against saved snapshots):
    python -m pytest core_engine/tests/test_report_baseline.py -q
"""
from __future__ import annotations

import json
import os
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest

# ── path setup (mirrors test_report_pipeline.py) ─────────────────────────────
_CORE = Path(__file__).resolve().parents[1]
_ROOT = _CORE.parent
for _p in (str(_CORE), str(_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

os.chdir(str(_CORE))

from adapters.asset import AssetValuationResult  # noqa: E402
from engines.base import AuditEntry              # noqa: E402
from reports.excel_builder import ExcelReportBuilder  # noqa: E402

# ── constants ─────────────────────────────────────────────────────────────────
_FIXTURES_DIR = Path(__file__).parent / "fixtures"
_FROZEN_DATE  = "2026-01-01"
_UPDATE       = os.environ.get("UPDATE_SNAPSHOTS", "0").strip() == "1"


# ── helpers ───────────────────────────────────────────────────────────────────

def _load_fixture(name: str) -> AssetValuationResult:
    """Load an input fixture JSON and construct AssetValuationResult."""
    raw = json.loads((_FIXTURES_DIR / name).read_text(encoding="utf-8"))
    return AssetValuationResult(
        asset_type        = raw["asset_type"],
        primary_purpose   = raw["primary_purpose"],
        primary_value     = Decimal(raw["primary_value"]) if raw.get("primary_value") else None,
        confidence        = raw["confidence"],
        weights_applied   = raw.get("weights_applied", {}),
        alternative_values= {k: Decimal(v) for k, v in raw.get("alternative_values", {}).items()},
        metadata          = raw.get("metadata", {}),
        audit_trail       = [AuditEntry(**e) for e in raw.get("audit_trail", [])],
        issues            = [],
        disclosures       = raw.get("disclosures", []),
    )


def _extract_cells(xlsx_path: str) -> dict[str, dict[str, Any]]:
    """Return {sheet_name: {cell_ref: value}} for every non-empty cell."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    result: dict[str, dict[str, Any]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        cells: dict[str, Any] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = cell.value
        if cells:
            result[sheet_name] = cells
    return result


def _build(result: AssetValuationResult | None,
           tmp_path: Path,
           name: str,
           style: str = "legacy") -> str:
    """Build a report with a frozen date. Returns output path."""
    builder = ExcelReportBuilder(result)
    builder.report_date = _FROZEN_DATE
    out = str(tmp_path / f"{name}.xlsx")
    builder.build(out, report_style=style)
    return out


def _snapshot_path(name: str) -> Path:
    return _FIXTURES_DIR / f"{name}.json"


def _load_snapshot(name: str) -> Any:
    return json.loads(_snapshot_path(name).read_text(encoding="utf-8"))


def _save_snapshot(name: str, data: Any) -> None:
    _snapshot_path(name).write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )


def _assert_or_generate(snapshot_name: str, actual: Any) -> None:
    """Save snapshot if UPDATE_SNAPSHOTS=1, otherwise assert equality."""
    if _UPDATE:
        _save_snapshot(snapshot_name, actual)
        return
    expected = _load_snapshot(snapshot_name)
    assert actual == expected, (
        f"Snapshot mismatch for '{snapshot_name}'.\n"
        f"Run with UPDATE_SNAPSHOTS=1 to regenerate if the change is intentional."
    )


# ── BL01 — residential legacy ─────────────────────────────────────────────────

def test_BL01_residential_legacy_baseline(tmp_path):
    """Residential result, legacy profile — all cell values match golden snapshot."""
    result = _load_fixture("report_residential.json")
    out    = _build(result, tmp_path, "residential_legacy", style="legacy")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_residential_legacy", cells)


# ── BL02 — commercial legacy ──────────────────────────────────────────────────

def test_BL02_commercial_legacy_baseline(tmp_path):
    """Commercial result, legacy profile — all cell values match golden snapshot."""
    result = _load_fixture("report_commercial.json")
    out    = _build(result, tmp_path, "commercial_legacy", style="legacy")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_commercial_legacy", cells)


# ── BL03 — residential detailed ───────────────────────────────────────────────

def test_BL03_residential_detailed_baseline(tmp_path):
    """Residential result, detailed profile — all cell values match golden snapshot."""
    result = _load_fixture("report_residential.json")
    out    = _build(result, tmp_path, "residential_detailed", style="detailed")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_residential_detailed", cells)


# ── BL04 — commercial detailed ────────────────────────────────────────────────

def test_BL04_commercial_detailed_baseline(tmp_path):
    """Commercial result, detailed profile — all cell values match golden snapshot."""
    result = _load_fixture("report_commercial.json")
    out    = _build(result, tmp_path, "commercial_detailed", style="detailed")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_commercial_detailed", cells)


# ── BL05 — None result legacy ─────────────────────────────────────────────────

def test_BL05_none_result_legacy_baseline(tmp_path):
    """PINS PRE-EXISTING IndexError when result=None — fix is out of scope for
    R1; this test must be updated when the bug is fixed.

    Root cause: ExcelReportBuilder.build() adds no sheets when result=None and
    no portfolio/composite kwargs are provided, so openpyxl raises IndexError
    ("At least one sheet must be visible") on save. The behaviour is pinned
    here so any silent regression or accidental fix is caught immediately."""
    with pytest.raises(IndexError, match="At least one sheet must be visible"):
        _build(None, tmp_path, "none_result_legacy", style="legacy")


# ── BL06 — sheet list (residential legacy) ───────────────────────────────────

def test_BL06_residential_sheet_list_baseline(tmp_path):
    """Sheet names for residential legacy must match golden snapshot."""
    import openpyxl
    result = _load_fixture("report_residential.json")
    out    = _build(result, tmp_path, "sheet_list_residential", style="legacy")
    wb     = openpyxl.load_workbook(out, data_only=True)
    sheet_list = wb.sheetnames
    _assert_or_generate("baseline_sheet_list_residential", sheet_list)


# ── BL07 — determinism proof ──────────────────────────────────────────────────

def test_BL07_two_runs_are_identical(tmp_path):
    """Two consecutive builds with the same frozen date must produce identical cells."""
    result = _load_fixture("report_residential.json")

    out1 = _build(result, tmp_path, "det_run1", style="legacy")
    out2 = _build(result, tmp_path, "det_run2", style="legacy")

    cells1 = _extract_cells(out1)
    cells2 = _extract_cells(out2)

    assert cells1 == cells2, "Two consecutive builds produced different cell values."


# ── BL08 — land legacy ────────────────────────────────────────────────────────

def test_BL08_land_legacy_baseline(tmp_path):
    """Land result, legacy profile — all cell values match golden snapshot."""
    result = _load_fixture("report_land.json")
    out    = _build(result, tmp_path, "land_legacy", style="legacy")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_land_legacy", cells)


# ── BL09 — land detailed ──────────────────────────────────────────────────────

def test_BL09_land_detailed_baseline(tmp_path):
    """Land result, detailed profile — all cell values match golden snapshot."""
    result = _load_fixture("report_land.json")
    out    = _build(result, tmp_path, "land_detailed", style="detailed")
    cells  = _extract_cells(out)
    _assert_or_generate("baseline_land_detailed", cells)
